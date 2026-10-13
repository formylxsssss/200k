#!/usr/bin/env python3
"""Build the tester-facing GUI operating manual."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


FONT_CN = "Microsoft YaHei"
FONT_CODE = "Consolas"
BLACK = "000000"
GRAY = "666666"
LIGHT_GRAY = "F3F5F7"
PALE_BLUE = "EEF4F8"
DARK_BLUE = "1F4E78"
LIGHT_BORDER = "D9D9D9"
RED = "B42318"


def set_run_font(run, size=10.5, bold=False, color=BLACK, name=FONT_CN):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def set_style_font(style, size, bold=False, color=BLACK, name=FONT_CN):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_BORDER, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, 8.5, color=GRAY)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, 8.5, color=GRAY)


class ManualBuilder:
    def __init__(self):
        self.doc = Document()
        self._configure_page()
        self._configure_styles()
        self._configure_header_footer()

    def _configure_page(self):
        section = self.doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.9)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)

    def _configure_styles(self):
        styles = self.doc.styles
        normal = styles["Normal"]
        set_style_font(normal, 10.5)
        normal.paragraph_format.line_spacing = 1.32
        normal.paragraph_format.space_after = Pt(6)

        for name, size, before, after in (
            ("Heading 1", 16, 18, 8),
            ("Heading 2", 13, 14, 6),
            ("Heading 3", 11.5, 10, 4),
        ):
            style = styles[name]
            set_style_font(style, size, bold=True, color=BLACK)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True

        title = styles["Title"]
        set_style_font(title, 26, bold=True, color=BLACK)
        title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(12)
        ppr = title.element.get_or_add_pPr()
        border = ppr.find(qn("w:pBdr"))
        if border is not None:
            ppr.remove(border)

        subtitle = styles.add_style("Manual Subtitle", WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(subtitle, 13, color=BLACK)
        subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(8)

        code = styles.add_style("Manual Code", WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(code, 9, color=BLACK, name=FONT_CODE)
        code.paragraph_format.left_indent = Cm(0.25)
        code.paragraph_format.right_indent = Cm(0.25)
        code.paragraph_format.space_before = Pt(3)
        code.paragraph_format.space_after = Pt(6)
        code.paragraph_format.line_spacing = 1.05

    def _configure_header_footer(self):
        section = self.doc.sections[0]
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = header.add_run("单通道水听器  GUI检测使用说明书")
        set_run_font(run, 8.5, color=GRAY)
        add_page_number(section.footer.paragraphs[0])

    def paragraph(self, text="", lead=None, lead_color=BLACK, keep=False):
        p = self.doc.add_paragraph()
        if keep:
            p.paragraph_format.keep_with_next = True
        if lead:
            run = p.add_run(lead)
            set_run_font(run, bold=True, color=lead_color)
        run = p.add_run(text)
        set_run_font(run)
        return p

    def heading(self, text, level=1):
        return self.doc.add_heading(text, level=level)

    def bullets(self, items):
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.3)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(item)
            set_run_font(run)

    def numbered(self, items):
        for index, item in enumerate(items, start=1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            p.paragraph_format.space_after = Pt(5)
            run = p.add_run(f"{index}. {item}")
            set_run_font(run)

    def code(self, text):
        p = self.doc.add_paragraph(style="Manual Code")
        p.paragraph_format.keep_together = True
        run = p.add_run(text)
        set_run_font(run, size=9, name=FONT_CODE)
        return p

    def warning(self, text):
        return self.paragraph(text, lead="警告：", lead_color=RED)

    def note(self, text):
        return self.paragraph(text, lead="注意：")

    def table(self, headers, rows, widths=None, font_size=9, center_cols=()):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table)
        hdr = table.rows[0]
        set_repeat_table_header(hdr)
        for idx, text in enumerate(headers):
            cell = hdr.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, DARK_BLUE)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(text))
            set_run_font(run, size=font_size, bold=True, color="FFFFFF")
            if widths:
                cell.width = Cm(widths[idx])

        for row_index, row_data in enumerate(rows):
            row = table.add_row()
            for idx, value in enumerate(row_data):
                cell = row.cells[idx]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                if row_index % 2 == 1:
                    set_cell_shading(cell, LIGHT_GRAY)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                run = p.add_run(str(value))
                set_run_font(run, size=font_size)
                if widths:
                    cell.width = Cm(widths[idx])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(1)
        return table

    def page_break(self):
        self.doc.add_page_break()

    def cover(self):
        section = self.doc.sections[0]
        section.different_first_page_header_footer = True
        for _ in range(4):
            self.doc.add_paragraph()
        self.doc.add_paragraph("单通道水听器", style="Title")
        self.doc.add_paragraph("GUI检测使用说明书", style="Title")
        p = self.doc.add_paragraph(style="Manual Subtitle")
        run = p.add_run("供生产检测  功能验证  现场联调人员使用")
        set_run_font(run, size=13, color=BLACK)
        for _ in range(3):
            self.doc.add_paragraph()
        self.table(
            ("项目", "内容"),
            (
                ("软件名称", "单通道水听器功能验证上位机"),
                ("主程序", "tools/host_gui/hydrophone_gui.py"),
                ("说明书版本", "V1.0"),
                ("软件基线日期", "2026-09-17"),
                ("适用人员", "生产检测、系统测试、现场联调和售后人员"),
            ),
            widths=(4.2, 10.4),
            font_size=10,
        )
        self.paragraph(
            "本说明书用于指导检测人员完成连接、自检、实时采集、自存储、任务表、低电平触发、SD卡管理和文件读取验证。操作前必须阅读安全注意事项。"
        )
        self.page_break()

    def build(self):
        self.heading("使用前必读", 1)
        self.paragraph(
            "本工具面向功能检测，不是数据长期归档软件。检测人员应先完成网络和设备状态检查，再按模式逐项测试。每次按钮操作都应等待成功或失败弹窗，不要连续重复点击。"
        )
        self.warning("格式化会永久删除目标卡中的全部数据。执行单卡或全部格式化前，必须确认卡槽、备份数据并取得测试负责人确认。")
        self.warning("格式化、写卡、目录读取、文件下载或SD扫描期间，严禁拔卡、断电或复位设备。")
        self.warning("实时采集和触发实时上传最高只能选择96 kSPS；128 kSPS和200 kSPS只用于写SD卡。")
        self.warning("设备复位或掉电后，授时、普通配置、任务表和触发配置均不会保留。重新上电后必须重新连接、授时和下发配置。")
        self.note("高采样率下逐样本保存CSV会明显增加电脑负载。96 kSPS检测应优先观察波形与统计，CSV只用于短时记录；长期原始数据应使用SD卡data.bin。")

        self.heading("快速操作顺序", 2)
        self.numbered((
            "设备上电，电脑网口设置为192.168.1.x/24，确保地址不占用192.168.1.80。",
            "运行hydrophone_gui.py，保持默认IP 192.168.1.80和端口5000、5001、5002。",
            "点击连接，确认顶部显示控制已连接、数据已连接。",
            "点击查询状态STATUS，再点击重新扫描SD_SCAN。",
            "需要准确记录时间、任务表或触发写卡时，点击上位机授时TIME。",
            "在对应模式区域设置该模式独立的采样率、精度和增益。",
            "点击启动按钮，等待配置成功和最终启动成功弹窗。",
            "检测结束后先点击对应停止按钮，确认设备回到空闲，再断开连接或拔卡。",
        ))

        self.heading("1 运行环境与连接", 1)
        self.heading("1.1 所需条件", 2)
        self.table(
            ("类别", "要求", "检测前确认"),
            (
                ("电脑", "Windows电脑，带可用有线网口", "□"),
                ("Python", "Python 3，能够使用tkinter、socket、threading标准库", "□"),
                ("设备", "单通道水听器设备已正确供电并完成固件烧录", "□"),
                ("网络", "电脑与设备直连或接入同一交换机", "□"),
                ("SD卡", "写卡测试使用可擦除测试卡，统一exFAT", "□"),
                ("信号源", "实时、增益和触发测试时准备稳定且不过载的输入信号", "□"),
            ),
            widths=(2.5, 10.3, 2.0),
            center_cols=(2,),
        )
        self.heading("1.2 启动程序", 2)
        self.numbered((
            "在PyCharm中打开tools/host_gui/hydrophone_gui.py。",
            "选择可运行Python 3的解释器。GUI本身不依赖第三方库。",
            "直接运行hydrophone_gui.py。窗口标题应显示单通道水听器功能验证上位机。",
            "如果左侧内容显示不全，把鼠标移到左侧操作区，使用滚轮或滚动条向下查看。",
        ))
        self.heading("1.3 网络设置", 2)
        self.table(
            ("项目", "默认值", "说明"),
            (
                ("设备IP", "192.168.1.80", "电脑不得使用相同地址"),
                ("子网掩码", "255.255.255.0", "电脑建议设置192.168.1.10等空闲地址"),
                ("控制端口", "5000", "命令和状态响应"),
                ("数据端口", "5001", "实时二进制采样数据"),
                ("文件端口", "5002", "SD目录浏览和文件下载"),
            ),
            widths=(3.0, 4.0, 7.8),
            center_cols=(1,),
        )
        self.note("点击连接后，GUI先连接数据端口，再连接控制端口。只有两个端口都成功时才会弹出连接完成。文件端口在读取目录或下载文件时按需连接。")

        self.heading("2 界面与通用规则", 1)
        self.heading("2.1 界面区域", 2)
        self.table(
            ("区域", "位置", "主要内容"),
            (
                ("连接设置", "顶部", "IP、三个端口、连接、断开、控制和数据连接状态"),
                ("时间和状态", "左侧上部", "电脑时间、设备授时、STATUS、SD状态和四卡槽信息"),
                ("四种模式", "左侧中部", "实时、自存储、任务表、低电平触发的独立参数与按钮"),
                ("SD操作", "左侧下部", "单卡/全部格式化、SD_DEMO写入测试"),
                ("文件读取", "左侧下部", "选择卡槽、读取目录、下载选中文件"),
                ("实时统计", "右侧顶部", "包数、样本数、实测SPS、丢包、序号、采样率和精度"),
                ("功能页签", "右侧主体", "实时波形、任务表、SD文件和日志"),
            ),
            widths=(3.0, 3.0, 8.8),
        )
        self.heading("2.2 参数范围", 2)
        self.table(
            ("参数", "可选值", "限制"),
            (
                ("实时采样率", "1/2/4/8/12/16/24/48/96 kSPS", "最高96 kSPS"),
                ("写卡采样率", "1/2/4/8/12/16/24/48/96/128/200 kSPS", "最高200 kSPS"),
                ("采样精度", "16 bit或24 bit", "界面不提供32 bit"),
                ("增益", "-30/-20/-10/0/+20/+40 dB", "默认0 dB"),
                ("触发持续时间", "0.001至86400秒", "最长24小时"),
                ("任务数量", "最多16项", "设备复位后清空"),
            ),
            widths=(3.2, 6.7, 4.9),
        )
        self.heading("2.3 模式互斥与按钮响应", 2)
        self.bullets((
            "实时、自存储、任务采集和触发采集不能同时运行。启动另一模式前，应先停止当前模式并查询STATUS确认空闲。",
            "实时和自存储的开始按钮会自动先下发本模式CONFIG，配置成功后再发送START。",
            "每次操作后等待弹窗。弹窗中的设备响应是成功或失败的最终依据，不能只看按钮是否按下。",
            "四种模式的采样率、精度和增益互相独立。修改一个模式不会同步修改其他模式。",
            "界面没有手动命令输入框。检测时只使用对应功能按钮。",
        ))

        self.heading("3 开机检查与授时", 1)
        self.heading("3.1 标准开机检查", 2)
        self.numbered((
            "点击连接并确认控制和数据均显示已连接。",
            "点击查询状态STATUS，检查设备模式、数据连接、采样缓冲、丢弃和漏采。",
            "点击重新扫描SD_SCAN，等待四个卡槽逐一检测完成。",
            "记录每个卡槽的R、U或未插卡状态。",
            "写卡测试前确认至少一张卡为可写。",
        ))
        self.table(
            ("显示", "含义", "处理"),
            (
                ("R", "卡已初始化、文件系统可用并可写", "可进行写卡测试"),
                ("U", "检测到卡，但文件系统不可用或挂载失败", "备份后格式化为exFAT，再扫描"),
                ("未插卡", "未检测到卡或底层初始化失败", "断电检查卡和卡槽，再重新扫描"),
                ("正在检测", "SD_SCAN尚未完成", "等待完成，不操作其他SD功能"),
            ),
            widths=(2.8, 6.6, 5.4),
        )
        self.heading("3.2 授时", 2)
        self.paragraph("时间设置默认自动跟随电脑时间。日期和时间会逐秒变化，点击上位机授时TIME时直接使用电脑当前时间。")
        self.numbered((
            "先确认电脑日期、时间和时区正确。",
            "保持时间来源显示电脑时间自动，或手动编辑年月日时分秒。",
            "手动修改任意时间字段后，自动更新会停止。",
            "点击上位机授时TIME并等待授时完成弹窗。",
            "需要重新跟随电脑时间时，点击恢复跟随电脑时间。",
        ))
        self.note("任务表依赖有效授时。自存储即使未授时也能开始，但目录和meta.txt中的日期时间不具备可追溯性。")

        self.heading("4 模式1 实时采集", 1)
        self.heading("4.1 操作步骤", 2)
        self.numbered((
            "确认顶部控制已连接、数据已连接。",
            "在模式1区域选择采样率、16/24 bit精度和增益。实时最高96 kSPS。",
            "需要保存短时CSV时，先选择文件，再勾选保存CSV。",
            "点击开始实时采集START_RT。等待配置成功后，GUI会自动发送启动命令。",
            "切换到实时波形页签，观察波形、Vpp和实时统计。",
            "测试结束点击停止实时采集STOP_RT，等待停止完成弹窗。",
        ))
        self.heading("4.2 合格判定", 2)
        self.table(
            ("检查项", "预期结果", "结果"),
            (
                ("启动响应", "弹出实时采集已启动，设备模式为实时传输", "□合格  □不合格"),
                ("实时数据", "数据包和样本数持续增长，波形不再显示暂无数据", "□合格  □不合格"),
                ("采样率", "实测SPS稳定接近包头显示的实际采样率", "□合格  □不合格"),
                ("连续性", "丢包计数不持续增加，设备DROP和MISS为0", "□合格  □不合格"),
                ("停止", "停止弹窗成功，STATUS回到空闲", "□合格  □不合格"),
            ),
            widths=(3.0, 8.2, 3.6),
        )
        self.note("实测SPS是电脑按短时间窗口统计的接收速率，刚启动、切换窗口或电脑负载较高时会短暂波动。最终采样率以实时包头和设备日志中的实际值为准。")
        self.warning("数据端口断开或GUI长期来不及读取时，设备可能自动停止实时采集。发现波形停止后先查询STATUS，不要直接反复点击START_RT。")

        self.heading("4.3 增益功能检测", 2)
        self.paragraph("使用幅度稳定、不会在+40 dB时造成过载的信号。保持采样率、精度、输入频率和幅度不变，依次测试六个增益档。")
        self.table(
            ("增益档", "相对0 dB名义幅度", "检测要点"),
            (
                ("-30 dB", "0.0316倍", "软件衰减，波形和SD数据均应生效"),
                ("-20 dB", "0.1倍", "PB9硬件支路"),
                ("-10 dB", "0.3162倍", "软件衰减，波形和SD数据均应生效"),
                ("0 dB", "1倍", "PB6硬件支路，默认档"),
                ("+20 dB", "10倍", "PB7硬件支路，防止ADC过载"),
                ("+40 dB", "100倍", "PB8硬件支路，输入必须足够小"),
            ),
            widths=(3.0, 4.5, 7.3),
            center_cols=(0, 1),
        )
        self.note("表中为名义比例，只用于快速功能判定。正式幅频和增益精度应使用校准设备并按硬件测试规范验收。")

        self.heading("5 模式2 自存储", 1)
        self.heading("5.1 操作步骤", 2)
        self.numbered((
            "停止其他采集，重新扫描SD卡，确认至少一张卡为可写。",
            "完成设备授时。",
            "在模式2区域选择采样率、精度和增益。写卡最高200 kSPS。",
            "点击开始自存储START_STORE，等待启动成功弹窗。",
            "点击STATUS，确认模式为自存储、REC=1并显示当前卡槽。",
            "达到预定时长后点击停止自存储STOP_STORE，等待文件关闭完成。",
            "确认STATUS回到空闲，再通过SD文件页签下载或断电拔卡检查文件。",
        ))
        self.note("自存储模式不向5001发送采样数据，因此右侧实时波形不会更新。这不是故障，应以STATUS、REC、SLOT和SD文件为准。")
        self.heading("5.2 文件检查", 2)
        self.code(
            "/HYDROPHONE/DATA/RUN_YYYYMMDD_HHMMSS_mmm_<RATE>Hz_MANUAL_<RUNID>/\n"
            "  meta.txt\n  data.bin\n  segment_end.txt\n  run_end.txt"
        )
        self.table(
            ("文件", "检查内容"),
            (
                ("meta.txt", "来源MANUAL、卡槽、采样率、实际采样率、精度、增益、开始时间"),
                ("data.bin", "文件大小大于0，且能按meta.txt中的精度解析"),
                ("segment_end.txt", "本卡段停止原因、样本数和字节数"),
                ("run_end.txt", "整次记录的最终样本数、字节数和停止统计"),
                ("switch_info.txt", "发生跨卡切换时记录上一卡和切换信息"),
            ),
            widths=(4.0, 10.8),
        )
        self.warning("停止响应返回前不要拔卡。异常断电可能导致最后一个文件未同步或目录损坏。")

        self.heading("6 模式3 任务表", 1)
        self.heading("6.1 添加任务", 2)
        self.numbered((
            "先授时，再扫描SD卡并确认有可写卡。",
            "在模式3区域选择该任务独立的采样率、精度和增益。",
            "设置开始年月日时分秒和结束年月日时分秒。结束必须晚于开始。",
            "设置间隔和次数。间隔表示相邻两次计划开始时间之间的时间。",
            "点击添加任务TASK_ADD，等待成功弹窗。任务添加成功后自动等待执行，不需要另按启动按钮。",
            "点击读取设备任务TASK_LIST，切换到右侧任务表确认任务内容。",
        ))
        self.note("任务表最多16项。次数最小为1。间隔小于单次持续时间时，下一次会在上一轮结束后尽快开始，不会并行采集。")
        self.heading("6.2 查询与删除", 2)
        self.bullets((
            "读取设备任务会读取0至15号表项并刷新右侧表格。",
            "删除单项前，先在右侧任务表中选中对应行，再点击删除选中任务。",
            "一键删除所有任务会停止正在运行的任务并清空全部表项，必须二次确认。",
            "任务完成后仍保留为已完成，直到被删除、清空或设备复位。",
            "设备复位或掉电后任务表丢失，重新上电后必须重新授时和添加。",
        ))
        self.heading("6.3 快速检测示例", 2)
        self.table(
            ("项目", "建议设置"),
            (
                ("开始时间", "当前电脑时间之后30秒"),
                ("结束时间", "开始后10秒"),
                ("间隔", "20秒"),
                ("次数", "2"),
                ("采样率", "1 kSPS或2 kSPS"),
                ("预期", "执行两次并产生两个TASK来源记录，表格执行次数更新"),
            ),
            widths=(4.0, 10.8),
        )

        self.heading("7 模式4 低电平触发", 1)
        self.heading("7.1 触发条件", 2)
        self.bullets((
            "触发输入固定为PC13低电平，不提供高电平触发选项。",
            "PC13低电平稳定至少10 ms后触发。高电平为未触发状态。",
            "触发后动作可选写入SD卡或实时上传。",
            "单次触发完成后回到空闲；连续触发完成后等待PC13恢复高电平，再重新布防。",
            "PC13一直保持低电平时不会重复触发；采集期间到来的新触发不会排队。",
        ))
        self.heading("7.2 操作步骤", 2)
        self.numbered((
            "确保设备空闲。触发写卡时先扫描并确认可写卡；触发实时上传时确认数据端口已连接。",
            "选择触发模式独立的采样率、精度和增益。实时上传最高96 kSPS。",
            "选择写入SD卡或实时上传，设置持续时间和单次/连续触发。",
            "点击下发触发配置TRIG_CONFIG，等待配置成功。",
            "确认PC13当前为高电平，再点击启动触发模式TRIG_ARM。",
            "点击查询触发状态TRIG_STATUS，确认设备处于等待低电平触发。",
            "将PC13拉低并保持至少10 ms，观察采集启动。",
            "连续触发测试时，采集结束后先释放为高电平，再产生下一次低电平。",
            "测试完成点击停止触发模式TRIG_STOP。",
        ))
        self.warning("如果TRIG_ARM时PC13已经为低电平，设备可能在10 ms消抖后立即触发。需要清晰边沿测试时，布防前必须确认PC13为高电平。")
        self.heading("7.3 合格判定", 2)
        self.table(
            ("检查项", "预期结果", "结果"),
            (
                ("高电平", "保持等待，不启动采集", "□合格  □不合格"),
                ("低电平", "稳定10 ms后按配置启动", "□合格  □不合格"),
                ("持续时间", "到达设定时长后自动停止本次采集", "□合格  □不合格"),
                ("单次", "完成后回到空闲", "□合格  □不合格"),
                ("连续", "释放为高后重新布防，下一次低电平可再次触发", "□合格  □不合格"),
                ("输出", "STORAGE生成TRIG目录；REALTIME显示HY实时波形", "□合格  □不合格"),
            ),
            widths=(3.0, 8.2, 3.6),
        )

        self.heading("8 SD卡管理", 1)
        self.heading("8.1 扫描和供电行为", 2)
        self.paragraph("SD_SCAN会逐槽上电、初始化、挂载并读取容量。扫描完成后四个卡槽全部断电；写卡时只给当前使用的卡保持供电，其他卡保持断电。插拔卡后必须重新扫描。")
        self.warning("不要在卡槽上电、扫描、格式化、写卡、LIST或GET期间插拔SD卡。安全做法是先停止采集，确认设备空闲，必要时设备断电后再换卡。")
        self.heading("8.2 单卡格式化", 2)
        self.numbered((
            "确认设备空闲，并备份目标卡数据。",
            "在SD卡操作区域选择卡槽1至4。",
            "点击格式化所选卡，核对警告中的卡槽。",
            "确认后等待按钮恢复和格式化完成弹窗。期间不要操作其他SD按钮。",
            "GUI会自动重新扫描，确认目标卡显示可写和容量正常。",
        ))
        self.heading("8.3 一键格式化全部卡", 2)
        self.warning("该按钮会清空所有当前已插入的SD卡，空卡槽只会跳过。只能使用专用测试卡执行。")
        self.numbered((
            "停止所有采集并确认设备空闲。",
            "确认所有已插入卡都允许擦除，完成必要备份。",
            "点击一键格式化全部卡，阅读危险操作提示并确认。",
            "等待最终弹窗。成功弹窗会显示成功格式化数量和跳过空卡槽数量。",
            "如果提示部分失败，记录设备返回内容，不要立即重试；先重新扫描并逐卡定位。",
        ))
        self.heading("8.4 SD_DEMO写入测试", 2)
        self.paragraph("SD_DEMO用于快速验证卡槽、挂载、建目录和写文件，不使用ADC，也不会删除已有文件。默认测试字节为4096。")
        self.numbered((
            "扫描SD卡并确认目标卡为可写。",
            "测试卡槽选择ALL或指定1至4。",
            "保持测试字节4096，或输入需要的正整数。",
            "点击写入假数据SD_DEMO，等待完成弹窗。",
            "在SD文件页签或读卡器中检查/HYDROPHONE/TEST/下的测试文件。",
        ))

        self.heading("9 通过网口读取SD文件", 1)
        self.paragraph("文件服务是只读功能，只支持目录浏览和下载，不支持上传、删除或重命名。设备必须处于空闲状态。")
        self.numbered((
            "停止实时、自存储和触发模式，查询STATUS确认空闲。",
            "在SD文件读取区域选择卡槽。当前目录为只读，默认/HYDROPHONE。",
            "点击读取目录LIST，右侧切换到SD文件页签。",
            "双击目录进入；使用返回上级返回；点击采集目录回到/HYDROPHONE。",
            "点击选择保存目录，指定电脑上的目标文件夹。",
            "选中文件，点击下载选中文件GET。等待下载完成弹窗。",
            "比较本地文件大小与列表中的SIZE；不一致的文件不得用于分析。",
        ))
        self.note("大文件下载期间窗口仍可响应，但不要同时启动采集、扫描、格式化或再次下载。网络中断后，删除或标记未完整下载的本地文件。")

        self.heading("10 波形 统计与CSV", 1)
        self.heading("10.1 实时统计", 2)
        self.table(
            ("字段", "含义", "检测提示"),
            (
                ("数据包", "本次实时会话收到的HY包数量", "应持续增加"),
                ("样本数", "已接收样本总数", "应持续增加"),
                ("实测SPS", "电脑短时间窗口统计的样本接收率", "稳态应接近实际采样率"),
                ("丢包", "根据实时包序号估算的缺失包数", "正常稳定测试应不持续增长"),
                ("序号", "最近一个实时包序号", "连续递增，允许uint32回绕"),
                ("采样率", "实时包头中的实际采样率", "分析时间轴使用此值"),
                ("精度", "16或24", "必须与配置一致"),
                ("最近码值/电压", "最新样本的码值和换算电压", "用于快速排查符号、偏置和幅度"),
            ),
            widths=(3.0, 6.5, 5.3),
        )
        self.heading("10.2 波形和电压", 2)
        self.paragraph("实时波形纵轴自动在mV和V之间切换，并显示最小值、最大值和Vpp。时间窗口可选择0.01至5秒。")
        self.code(
            "16 bit: V = code / 32768 x 4.096 V\n"
            "24 bit: V = code / 8388608 x 4.096 V\n"
            "Vpp = Vmax - Vmin"
        )
        self.note("该电压是ADS127L11差分输入AN1+减AN1-的码值换算，不等同于水听器声压，也未自动包含传感器灵敏度、前端误差和系统标定系数。")
        self.heading("10.3 CSV保存", 2)
        self.bullets((
            "先点击选择文件，再勾选保存CSV；如果未选择文件，程序会在当前目录生成默认文件名。",
            "CSV列包括电脑时间ms、包序号、包内索引、原始码值和电压V。",
            "CSV中的电脑时间是收包时刻，不是每个样本的精确设备时间。",
            "停止采集后取消勾选保存CSV，确保文件关闭后再打开或复制。",
            "高采样率长时间CSV会占用大量CPU和磁盘，不作为无丢样验收的唯一依据。",
        ))

        self.heading("11 离线BIN分析", 1)
        self.paragraph("从SD卡下载或通过读卡器复制data.bin后，可运行tools/bin_analyzer/bin_analyzer.py进行时域和dBV底噪分析。")
        self.numbered((
            "把data.bin与同一运行目录中的meta.txt放在同一文件夹。",
            "在tools/bin_analyzer目录安装requirements.txt中的依赖。",
            "运行bin_analyzer.py并选择data.bin。",
            "确认程序读取到实际采样率、16/24 bit精度和增益。",
            "按需要设置分析起始时间、时长和频带，查看时域、Vpp、RMS和dBV/√Hz。",
        ))
        self.note("跨卡记录需要按meta.txt中的segment顺序合并。不同卡上的同名RUN目录属于同一次记录，不要按文件复制时间排序。")

        self.heading("12 常见问题与处理", 1)
        self.table(
            ("现象", "优先检查", "处理方法"),
            (
                ("连接失败或WinError 10054", "网段、IP冲突、设备复位、旧客户端占用", "关闭旧客户端，检查192.168.1.x/24，重新连接"),
                ("按钮点击后无成功弹窗", "控制连接、日志、命令超时", "不要连续点击；查询STATUS，必要时断开后重连"),
                ("START_RT失败", "数据端口是否已连接", "先确保数据已连接，再启动实时采集"),
                ("实时波形卡顿", "CSV、电脑负载、波形窗口、采样率", "关闭CSV，缩短时间窗口，关闭其他高负载程序"),
                ("丢包持续增加", "接收线程、网线、电脑性能", "停止采集，检查网络和负载，降低采样率后复测"),
                ("自存储时没有波形", "当前模式是否为写卡", "正常现象，用STATUS和SD文件检查"),
                ("SD显示U", "文件系统不可用", "备份数据后格式化为exFAT并重新扫描"),
                ("SD显示未插卡", "卡未识别或初始化失败", "设备空闲或断电后重插，换卡/换槽，再扫描"),
                ("START_STORE失败", "是否至少有一张可写卡", "先SD_SCAN，确认至少一个卡槽为R"),
                ("任务不执行", "是否授时、开始时间、卡状态、任务是否重叠", "重新授时，读取任务并核对时间"),
                ("触发不动作", "PC13电平、布防状态、持续低电平时间", "TRIG_STATUS确认pin和mode，低电平保持至少10 ms"),
                ("LIST或GET失败", "设备是否空闲、卡是否可写/可挂载", "停止所有模式，确认空闲后重新读取"),
                ("文件时间不正确", "是否在采集前授时", "校准电脑时间，重新TIME后再采集"),
                ("增益变化不明显", "信号幅度、频率、过载、硬件支路", "降低输入，固定频率，记录码值/Vpp并分档比较"),
            ),
            widths=(4.0, 5.0, 5.8),
            font_size=8.6,
        )

        self.heading("13 标准功能检测记录", 1)
        self.paragraph("建议每块设备至少完成以下项目。结果栏由检测人员填写，失败项目应记录日志、设备响应、复现条件和对应文件。")
        self.table(
            ("序号", "检测项目", "建议条件", "预期结果", "结果"),
            (
                ("1", "连接与STATUS", "默认IP和端口", "控制/数据连接成功，STATUS可读", "□通过 □失败"),
                ("2", "四卡槽扫描", "逐槽插入测试卡", "R/U/未插卡与实物一致", "□通过 □失败"),
                ("3", "授时", "电脑时间自动", "TIME成功，记录目录时间正确", "□通过 □失败"),
                ("4", "实时16 bit", "1 kSPS，0 dB", "波形、SPS、丢包正常", "□通过 □失败"),
                ("5", "实时24 bit", "96 kSPS，0 dB", "持续接收，设备DROP/MISS为0", "□通过 □失败"),
                ("6", "六档增益", "稳定小信号", "各档幅度方向和名义比例合理", "□通过 □失败"),
                ("7", "自存储", "200 kSPS，24 bit，短时", "正常停止并生成完整记录文件", "□通过 □失败"),
                ("8", "每槽SD_DEMO", "1至4号卡各4096字节", "对应TEST文件存在且大小正确", "□通过 □失败"),
                ("9", "任务表", "2次短任务", "按时执行、可查询、可单删和全清", "□通过 □失败"),
                ("10", "单次触发写卡", "PC13低电平", "触发一次并生成TRIG记录", "□通过 □失败"),
                ("11", "连续触发实时", "两次高低电平循环", "两段实时数据均正常", "□通过 □失败"),
                ("12", "文件LIST/GET", "下载meta和data.bin", "本地大小与SIZE一致", "□通过 □失败"),
                ("13", "单卡格式化", "专用空白测试卡", "完成弹窗且重新扫描为R", "□通过 □失败"),
                ("14", "全部卡格式化", "仅限可擦除测试卡", "formatted/skipped/failed数量正确", "□通过 □失败"),
                ("15", "复位恢复", "设备复位后重连", "识别断线，时间/任务需重新配置", "□通过 □失败"),
            ),
            widths=(1.2, 3.5, 3.9, 4.2, 2.2),
            font_size=8.2,
            center_cols=(0, 4),
        )

        self.heading("14 检测信息记录", 1)
        self.table(
            ("项目", "填写内容"),
            (
                ("设备编号", ""),
                ("固件版本或烧录文件", ""),
                ("GUI版本或文件日期", ""),
                ("电脑IP和系统", ""),
                ("SD卡品牌 容量 卡槽", ""),
                ("信号源型号和设置", ""),
                ("检测日期", ""),
                ("检测人员", ""),
                ("异常说明和附件路径", ""),
                ("复核人员", ""),
            ),
            widths=(5.2, 9.6),
        )

        self.heading("附录 默认参数与文件来源", 1)
        self.table(
            ("项目", "默认值或规则"),
            (
                ("设备地址", "192.168.1.80"),
                ("端口", "控制5000  数据5001  文件5002"),
                ("实时模式默认", "96 kSPS  24 bit  0 dB"),
                ("其他模式默认", "200 kSPS  24 bit  0 dB"),
                ("触发默认", "写入SD卡  连续触发  持续10秒"),
                ("波形窗口默认", "0.2秒，可选0.01至5秒"),
                ("SD文件系统", "exFAT"),
                ("SD选择顺序", "卡槽1至4顺位使用"),
                ("程序基线", "tools/host_gui/hydrophone_gui.py"),
                ("离线分析", "tools/bin_analyzer/bin_analyzer.py"),
            ),
            widths=(5.0, 9.8),
        )

    def save(self, output):
        output.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("output") / "单通道水听器GUI检测使用说明书_V1.0.docx",
    )
    args = parser.parse_args()
    builder = ManualBuilder()
    builder.cover()
    builder.build()
    builder.save(args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
