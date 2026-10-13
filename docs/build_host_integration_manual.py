#!/usr/bin/env python3
"""Generate the host integration manual for the current firmware protocol."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

FONT_ASCII = "Calibri"
FONT_CJK = "Microsoft YaHei"
FONT_CODE = "Consolas"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "172B3A"
MUTED = "66727D"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CODE_FILL = "F6F8FA"
NOTE_FILL = "EEF5FA"
WARN_FILL = "FFF4CE"
WHITE = "FFFFFF"
BORDER = "B8C4CE"


def set_run_font(run, *, size=None, bold=None, italic=None, color=None, code=False):
    name = FONT_CODE if code else FONT_ASCII
    east_asia = FONT_CODE if code else FONT_CJK
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, size, color="000000", bold=False, name=FONT_ASCII, east_asia=FONT_CJK):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa):
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {CONTENT_WIDTH_DXA}: {widths_dxa}")

    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440.0)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def mark_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_paragraph_shading(paragraph, fill, left_border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if left_border:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), left_border)
        p_bdr.append(left)


def create_numbering(doc, fmt, text, left=540, hanging=270):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(start)
    lvl.append(num_fmt)
    lvl.append(lvl_text)
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(tabs)
    p_pr.append(ind)
    p_pr.append(spacing)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT_ASCII)
    r_fonts.set(qn("w:hAnsi"), FONT_ASCII)
    r_fonts.set(qn("w:eastAsia"), FONT_CJK)
    r_pr.append(r_fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return abstract_id, num_id


def create_num_instance(doc, abstract_id):
    numbering = doc.part.numbering_part.element
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    lvl_override = OxmlElement("w:lvlOverride")
    lvl_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    lvl_override.append(start_override)
    num.append(lvl_override)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(result)
    run._r.append(end)
    return run


class ManualBuilder:
    def __init__(self):
        self.doc = Document()
        self._configure_document()
        self.bullet_abstract, self.bullet_num = create_numbering(self.doc, "bullet", "•")
        self.decimal_abstract, self.decimal_num = create_numbering(self.doc, "decimal", "%1.")

    def _configure_document(self):
        doc = self.doc
        section = doc.sections[0]
        section.page_width = Inches(PAGE_WIDTH_DXA / 1440.0)
        section.page_height = Inches(PAGE_HEIGHT_DXA / 1440.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        section.different_first_page_header_footer = True

        styles = doc.styles
        normal = styles["Normal"]
        set_style_font(normal, 11)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.25

        for name, size, color, before, after in (
            ("Heading 1", 16, BLUE, 18, 10),
            ("Heading 2", 13, BLUE, 14, 7),
            ("Heading 3", 12, DARK_BLUE, 10, 5),
        ):
            style = styles[name]
            set_style_font(style, size, color=color, bold=True)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.line_spacing = 1.0
            style.paragraph_format.keep_with_next = True

        title = styles["Title"]
        set_style_font(title, 27, color="000000", bold=True)
        title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(10)
        title_ppr = title.element.get_or_add_pPr()
        title_border = title_ppr.find(qn("w:pBdr"))
        if title_border is not None:
            title_ppr.remove(title_border)

        title = styles.add_style("Manual Cover Title", WD_STYLE_TYPE.PARAGRAPH)
        title.base_style = styles["Title"]
        set_style_font(title, 27, color="000000", bold=True)
        title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(10)

        subtitle = styles.add_style("Manual Cover Subtitle", WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(subtitle, 14, color=DARK_BLUE)
        subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(8)

        code = styles.add_style("Manual Code", WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(code, 9, color="263238", name=FONT_CODE, east_asia=FONT_CODE)
        code.paragraph_format.left_indent = Inches(0.08)
        code.paragraph_format.right_indent = Inches(0.08)
        code.paragraph_format.space_before = Pt(4)
        code.paragraph_format.space_after = Pt(8)
        code.paragraph_format.line_spacing = 1.0

        caption = styles.add_style("Manual Caption", WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(caption, 9, color=MUTED, bold=True)
        caption.paragraph_format.space_before = Pt(4)
        caption.paragraph_format.space_after = Pt(4)
        caption.paragraph_format.keep_with_next = True

        self._configure_header_footer(section)

        settings = doc.settings.element
        update_fields = OxmlElement("w:updateFields")
        update_fields.set(qn("w:val"), "true")
        settings.append(update_fields)

        core = doc.core_properties
        core.title = "单通道自容式水听器上位机对接手册"
        core.subject = "TCP 控制、实时数据、任务、触发、SD 卡及文件服务协议"
        core.author = "项目开发组"
        core.keywords = "水听器, 上位机, TCP, ADS127L11, W5500, SD"
        core.comments = "根据 single_channel_device 当前源码生成"

    def _configure_header_footer(self, section):
        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run("单通道自容式水听器  |  上位机对接手册")
        set_run_font(run, size=8.5, color=MUTED)

        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run("第 ")
        set_run_font(run, size=8.5, color=MUTED)
        field = add_field(p, " PAGE ", "1")
        set_run_font(field, size=8.5, color=MUTED)
        run = p.add_run(" 页")
        set_run_font(run, size=8.5, color=MUTED)

        first_header = section.first_page_header
        first_header.paragraphs[0].text = ""
        first_footer = section.first_page_footer
        p = first_footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("文档版本 V1.2  |  固件基线 2026-09-17")
        set_run_font(run, size=8.5, color=MUTED)

    def paragraph(self, text="", *, bold_lead=None, style=None, keep=False):
        p = self.doc.add_paragraph(style=style)
        if bold_lead and text.startswith(bold_lead):
            first = p.add_run(bold_lead)
            set_run_font(first, bold=True)
            rest = p.add_run(text[len(bold_lead):])
            set_run_font(rest)
        else:
            run = p.add_run(text)
            set_run_font(run)
        p.paragraph_format.keep_together = keep
        return p

    def heading(self, text, level=1):
        p = self.doc.add_paragraph(style=f"Heading {level}")
        run = p.add_run(text)
        set_run_font(run, size={1: 16, 2: 13, 3: 12}[level], bold=True,
                     color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level])
        return p

    def bullets(self, items):
        for item in items:
            p = self.doc.add_paragraph()
            apply_num(p, self.bullet_num)
            run = p.add_run(item)
            set_run_font(run)

    def numbered(self, items):
        num_id = create_num_instance(self.doc, self.decimal_abstract)
        for item in items:
            p = self.doc.add_paragraph()
            apply_num(p, num_id)
            run = p.add_run(item)
            set_run_font(run)

    def code(self, text):
        p = self.doc.add_paragraph(style="Manual Code")
        p.paragraph_format.keep_together = True
        set_paragraph_shading(p, CODE_FILL)
        lines = text.splitlines() or [""]
        for index, line in enumerate(lines):
            if index:
                p.add_run().add_break()
            run = p.add_run(line)
            set_run_font(run, size=9, code=True, color="263238")
        return p

    def note(self, label, text, warning=False):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.08)
        p.paragraph_format.right_indent = Inches(0.08)
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.2
        set_paragraph_shading(p, WARN_FILL if warning else NOTE_FILL, "C69214" if warning else BLUE)
        lead = p.add_run(f"{label}：")
        set_run_font(lead, bold=True, color="6B5200" if warning else DARK_BLUE)
        body = p.add_run(text)
        set_run_font(body)
        return p

    def table(self, headers, rows, widths_dxa, *, code_columns=(), small=False):
        table = self.doc.add_table(rows=1, cols=len(headers))
        set_table_geometry(table, widths_dxa)
        set_table_borders(table)
        table.rows[0]._tr.get_or_add_trPr()
        mark_repeat_table_header(table.rows[0])
        for index, header in enumerate(headers):
            cell = table.rows[0].cells[index]
            set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(str(header))
            set_run_font(run, size=9 if small else 9.5, bold=True, color=INK)
        for row_index, row in enumerate(rows):
            table_row = table.add_row()
            table_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            cells = table_row.cells
            for col_index, value in enumerate(row):
                cell = cells[col_index]
                if row_index % 2 == 1:
                    set_cell_shading(cell, "FAFBFC")
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                run = p.add_run(str(value))
                set_run_font(run, size=8.5 if small else 9.2, code=col_index in code_columns)
        after = self.doc.add_paragraph()
        after.paragraph_format.space_after = Pt(2)
        return table

    def caption(self, text):
        p = self.doc.add_paragraph(style="Manual Caption")
        run = p.add_run(text)
        set_run_font(run, size=9, bold=True, color=MUTED)
        return p

    def page_break(self):
        self.doc.add_page_break()

    def cover(self):
        for _ in range(4):
            self.doc.add_paragraph()
        kicker = self.doc.add_paragraph()
        kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
        kicker.paragraph_format.space_after = Pt(18)
        run = kicker.add_run("TECHNICAL INTEGRATION GUIDE")
        set_run_font(run, size=10, bold=True, color=BLUE)

        p = self.doc.add_paragraph(style="Title")
        run = p.add_run("单通道自容式水听器")
        set_run_font(run, size=27, bold=True, color="000000")
        p = self.doc.add_paragraph(style="Manual Cover Title")
        p.paragraph_format.space_after = Pt(16)
        run = p.add_run("上位机对接手册")
        set_run_font(run, size=27, bold=True, color="000000")

        p = self.doc.add_paragraph(style="Manual Cover Subtitle")
        run = p.add_run("TCP 控制 · 实时数据 · 自存储 · 任务表 · 低电平触发 · SD 文件服务")
        set_run_font(run, size=13.5, color=DARK_BLUE)

        for _ in range(3):
            self.doc.add_paragraph()

        meta = [
            ("文档版本", "V1.2"),
            ("固件工程", "single_channel_device"),
            ("参考上位机", "tools/host_gui/hydrophone_gui.py"),
            ("代码基线日期", "2026-09-17"),
            ("适用对象", "上位机开发、系统测试、现场联调人员"),
        ]
        table = self.doc.add_table(rows=len(meta), cols=2)
        set_table_geometry(table, [2340, 7020])
        for row, (label, value) in zip(table.rows, meta):
            set_cell_shading(row.cells[0], LIGHT_BLUE)
            for cell in row.cells:
                set_cell_margins(cell, top=100, bottom=100)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = row.cells[0].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(label)
            set_run_font(r, size=10, bold=True, color=DARK_BLUE)
            p = row.cells[1].paragraphs[0]
            r = p.add_run(value)
            set_run_font(r, size=10)
        # Cover metadata intentionally uses a quiet, borderless grid.
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), "nil")
            borders.append(element)
        tbl_pr.append(borders)
        self.page_break()

    def front_matter(self):
        self.heading("文档说明", 1)
        self.note(
            "适用范围",
            "本手册描述当前工作区中 single_channel_device 固件与 hydrophone_gui.py 的实际接口。目标规划中尚未落地的能力不作为已实现功能写入。",
        )
        self.paragraph(
            "设备以 W5500 提供三个独立 TCP 服务：控制命令、实时采样数据和 SD 卡只读文件服务。上位机需要分别维护连接，并根据工作模式处理文本响应或二进制流。"
        )
        self.note(
            "版本原则",
            "固件命令、返回字段、状态机或数据包发生变化时，应同步更新本手册。本文所述容量和 200 kS/s 为代码设计上限，最终交付前仍需在目标硬件、网口链路和实际 SD 卡上完成压力验证。",
            warning=True,
        )
        self.heading("V1.2 对接基线", 2)
        self.bullets([
            "上位机产品界面只开放 16 bit 和 24 bit；固件中的 32 bit 仅为兼容分支，不属于本版对外验收范围。",
            "普通实时与触发实时最高 96 kS/s；128 kS/s、200 kS/s 仅用于立即自存储、任务表或触发写卡。",
            "新增一键格式化全部已插入卡：SD_FORMAT SLOT=ALL FS=EXFAT CONFIRM=ERASE_ALL；空卡槽跳过，必须强确认。",
            "连续触发实时流的 sequence 在设备本次上电期间保持单调递增，不因每段触发重新清零。",
            "配置、授时、任务表和触发配置均在 RAM 中，设备复位或掉电后不会恢复；上位机重连后必须重新同步。",
        ])

        self.heading("目录", 1)
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        add_field(p, ' TOC \\o "1-1" \\h \\z \\u ', "请在 Word 中更新目录")
        self.page_break()

    def build_content(self):
        self.section_overview()
        self.section_function_recipes()
        self.section_transport()
        self.section_config()
        self.section_status_time()
        self.section_realtime()
        self.section_manual_storage()
        self.section_tasks()
        self.section_trigger()
        self.section_sd()
        self.section_file_service()
        self.section_offline_data()
        self.section_concurrency()
        self.section_troubleshooting()
        self.section_acceptance()
        self.section_delivery_requirements()
        self.appendix_commands()
        self.appendix_limits()

    def recipe(self, title, purpose, syntax, exchange, implementation):
        self.heading(title, 2)
        self.paragraph("功能：" + purpose, bold_lead="功能：")
        self.paragraph("指令格式：", bold_lead="指令格式：")
        self.code(syntax)
        example_label = self.paragraph("交互示例（> 为上位机发送，< 为设备返回）：")
        example_label.paragraph_format.keep_with_next = True
        self.code(exchange)
        self.paragraph("实现要点：" + implementation, bold_lead="实现要点：")

    def section_function_recipes(self):
        self.heading("逐功能对接实例", 1)
        self.note(
            "阅读方法",
            "本章按客户界面上的功能组织。除特别注明的文件服务外，所有指令发送到 TCP 5000；每条指令末尾追加 ASCII LF (\\n)。"
            "示例中的 SEQ=<n> 由设备分配，不能照抄为固定序号；为便于阅读而折行的长命令和长响应，实际都必须是一整行。"
            "响应正文以实际收到的内容为准。后续章节给出字段表和二进制格式。",
        )

        self.recipe(
            "功能 01：连接与设备状态查询",
            "确认控制连接、当前模式、数据端连接、卡数量与采集缓冲状态。",
            "STATUS\\n",
            "> STATUS\n< SEQ=<n> STATUS=0 MSG=IDLE CTRL=1 DATA=0 SDW=1 REC=0\n"
            "  SLOT=0 RUN=0 SAMPLES=0 TRIG_PIN=1 TOUT=STORAGE TDUR=10000\n"
            "  BUF=0 DROP=0 MISS=0 SW=0",
            "先连接 5000，再发送 STATUS；按空格拆出 MSG 后的模式和 KEY=VALUE。"
            "不要把 SLOT=0 或 SDW=1 当成采集已启动；以模式和 REC 为准。"
            "状态响应可能较长，控制端必须按 LF 收完整行。",
        )

        self.recipe(
            "功能 02：设置可变采样率",
            "设置下一次普通实时或立即自存储采集的目标采样率。只接受 1/2/4/8/12/16/24/48/96/128/200 kS/s 离散挡位。",
            "CONFIG RATE=<1000|2000|4000|8000|12000|16000|24000|48000|96000|128000|200000> PREC=<16|24> GAIN=<dB>\\n",
            "> CONFIG RATE=96000 PREC=24 GAIN=0\n"
            "< SEQ=<n> STATUS=0 MSG=CONFIG_OK rate=96000 prec=24BIT\n"
            "  gain=+0DB hw=+0DB sw=+0DB branch=0x01",
            "RATE 单位是样本/秒，不是 ADC 时钟频率。配置返回的 rate 是请求目标，不是实测值。"
            "每次均发送 RATE、PREC、GAIN 三个字段；省略字段会回到默认配置，不表示沿用上次值。"
            "修改运行中的采样率必须先 STOP_RT/STOP_STORE，重新 CONFIG 后再 START。"
            "实时模式读每个 HY 包头的 sample_rate；存储模式读 meta.txt 的 actual_sample_rate_hz。",
        )
        self.table(
            ["操作", "指令", "生效时机", "核验位置"],
            [
                ("改为 1 kS/s", "CONFIG RATE=1000 PREC=24 GAIN=0", "下次 START_RT/START_STORE", "HY 包头或 meta"),
                ("改为 200 kS/s", "CONFIG RATE=200000 PREC=24 GAIN=0", "仅下次存储启动", "meta"),
                ("改为 96 kS/s", "CONFIG RATE=96000 PREC=24 GAIN=0", "下次 START_RT/START_STORE", "HY 包头或 meta"),
                ("改为 128 kS/s", "CONFIG RATE=128000 PREC=24 GAIN=0", "仅存储模式", "meta"),
            ],
            [1600, 4140, 2050, 1570],
            code_columns=(1,),
            small=True,
        )
        self.note(
            "实际采样率",
            "定时器把目标采样率乘以 64 得到目标 ADC 时钟，再按整数周期分频；因此某些目标值会量化到邻近实际值。"
            "内部 HSI 时钟还存在频率误差。实时上传最高使用 96 kS/s，128/200 kS/s 用于存储模式；交付前仍需在实际 SD 卡上完成压力验证。",
            warning=True,
        )

        self.recipe(
            "功能 03：切换采样精度与字节宽度",
            "为下一次普通采集指定每样本 16 或 24 位。",
            "CONFIG RATE=<Hz> PREC=<16|24> GAIN=<dB>\\n",
            "> CONFIG RATE=200000 PREC=16 GAIN=0\n"
            "< SEQ=<n> STATUS=0 MSG=CONFIG_OK rate=200000 prec=16BIT\n"
            "  gain=+0DB hw=+0DB sw=+0DB branch=0x01",
            "16 位让 ADC 输出 2 字节并按有符号 int16 小端传输；24 位为 3 字节有符号小端。"
            "第三方上位机不得在产品界面提供 32 bit 选项；固件虽保留 24 位结果符号扩展为 int32 的兼容分支，但不作为本版协议承诺。"
            "实时包头 precision 和 meta.txt 的 bytes_per_sample 必须与解码器一致。",
        )
        self.code(
            "PREC=16 -> 2 bytes/sample -> signed int16 LE\n"
            "PREC=24 -> 3 bytes/sample -> signed int24 LE"
        )

        self.recipe(
            "功能 04：设置硬件或软件增益",
            "在固件预定义的六档增益中选择，默认 0 dB。",
            "CONFIG RATE=<Hz> PREC=<bits> GAIN=<-30|-20|-10|0|20|40>\\n",
            "> CONFIG RATE=200000 PREC=24 GAIN=40\n"
            "< SEQ=<n> STATUS=0 MSG=CONFIG_OK rate=200000 prec=24BIT\n"
            "  gain=+40DB hw=+40DB sw=+0DB branch=0x04",
            "硬件 -20/0/+20/+40 dB 分别使用 PB9/PB6/PB7/PB8；软件 -10/-30 dB 固定使用 PB6 的 0 dB 支路，"
            "在样本进入实时发送和 SD 存储公共缓冲区前乘以数字系数。PC4/PC5 不参与控制。对应关系见 3.2 节。",
        )
        self.note(
            "参数校验",
            "RATE 不在离散挡位、PREC 不是 16/24、GAIN 不在六档之一时，上位机应在本地直接拒绝；固件解析失败时可能直接丢弃该行，"
            "此时不是 STATUS=-1，而是控制口超时。上位机应先做本地校验，超时后不要盲目重复启动命令。",
            warning=True,
        )

        self.recipe(
            "功能 05：上位机年月日时分秒授时",
            "让任务时间和记录元数据使用绝对 Unix 毫秒。",
            "TIME <unix_epoch_ms>\\n",
            "> TIME 1789351200000\n< SEQ=<n> STATUS=0 MSG=TIME_OK",
            "示例对应北京时间 2026-09-14 10:00:00。GUI 使用带时区的日期时间转换后发送整数毫秒；"
            "不要把秒值误当毫秒。固件只确认接收，不校验日期合理性；上位机必须自己检查时间、时区和电脑时钟。"
            "重启后需重新授时，任务表也会清空。",
        )
        self.code(
            "from datetime import datetime, timezone, timedelta\n"
            "cst = timezone(timedelta(hours=8))\n"
            "dt = datetime(2026, 9, 14, 10, 0, 0, tzinfo=cst)\n"
            "unix_ms = int(dt.timestamp() * 1000)  # 1789351200000"
        )

        self.recipe(
            "功能 06：启动和停止实时采集",
            "把 ADC 样本持续发送给上位机，不写 SD。",
            "START_RT\\n\nSTOP_RT\\n",
            "> CONFIG RATE=96000 PREC=24 GAIN=0\n"
            "< SEQ=<n> STATUS=0 MSG=CONFIG_OK ...\n"
            "> START_RT\n< SEQ=<n> STATUS=0 MSG=START_REALTIME\n"
            "> STOP_RT\n< SEQ=<n> STATUS=0 MSG=STOP_REALTIME",
            "先连接 TCP 5001 并持续接收，再连接/使用 5000 下发 START_RT。5001 上是 HY 二进制流，"
            "不能按文本或每次 recv 恰好一包来解析；每包 16 字节头 + payload_len。"
            "若 5001 断开或缓冲首次溢出，设备会自动停止，需重连并重新 START_RT。",
        )
        self.code(
            "import struct\n"
            "header = recv_exact(data_socket, 16)\n"
            "magic, ver, prec, seq, rate, count, size = (\n"
            "    struct.unpack('<2sBBIIHH', header))\n"
            "assert magic == b'HY' and ver == 1\n"
            "assert size == count * {16: 2, 24: 3}[prec]\n"
            "payload = recv_exact(data_socket, size)"
        )
        self.paragraph("上述 recv_exact 必须循环接收直到字节数满足要求；字段说明及更多拆包规则见 5.2～5.6 节。")

        self.recipe(
            "功能 07：扫描四卡并查看剩余空间",
            "识别 1～4 号卡槽是否插卡、是否有可用文件系统，以及当前缓存的空闲空间。",
            "SD_SCAN\\n\nSD_STATUS\\n",
            "> SD_SCAN\n"
            "< SEQ=<n> STATUS=0 MSG=SD scan=1 present=2 writable=1\n"
            "  total_mb=<MB> S1=R:<freeMB>MB S2=U S3=- S4=-\n"
            "> SD_STATUS\n< SEQ=<n> STATUS=0 MSG=SD scan=1 ...",
            "SD_SCAN 真正逐槽上电、初始化和读容量，只有 IDLE 可执行；忙时返回 SD_SCAN_BUSY。"
            "SD_STATUS 只读最近扫描缓存，可在采集中调用，但插拔卡后不会自动刷新。"
            "R 表示文件系统可挂载，U 表示检测到卡但不可用，- 可能是未插卡或初始化/容量读取失败，不能绝对当成没卡。",
        )

        self.recipe(
            "功能 08：立即自存储与手动停止",
            "从当前时刻开始写 SD，按物理卡槽 1→4 选择可写卡，必要时自动跨卡。",
            "START_STORE\\n\nSTOP_STORE\\n",
            "> SD_SCAN\n< SEQ=<n> STATUS=0 MSG=SD scan=1 present=2 writable=2 ...\n"
            "> CONFIG RATE=200000 PREC=24 GAIN=0\n"
            "< SEQ=<n> STATUS=0 MSG=CONFIG_OK ...\n"
            "> START_STORE\n< SEQ=<n> STATUS=0 MSG=START_STORAGE\n"
            "> STATUS\n< SEQ=<n> STATUS=0 MSG=STORAGE_IMMEDIATE ... REC=1 SLOT=1 ...\n"
            "> STOP_STORE\n< SEQ=<n> STATUS=0 MSG=STOP_STORAGE",
            "启动前设备会扫描并缓存可写卡队列、打开首张卡，然后才启动 ADC。"
            "停止时先停采样，再排空缓冲并关闭文件；等待 STOP_STORE 返回并确认 STATUS 为 IDLE、REC=0 后，再检查 DROP/MISS 与记录文件。"
            "控制回复不能代替文件完整性校验；异常断电或写卡错误后，应先在电脑上检查文件再拔换或复用卡。"
            "无卡或全卡不可写时 START_STORE 返回 STATUS=-1，MSG 为 SD 摘要。",
        )
        self.note(
            "跨卡完整性",
            "有后续卡且当前卡估算剩余约 16 MiB 时提前切换。换卡期间 640 KiB 采样缓冲约可承受 819 ms/200 kS/s；"
            "这只是设计余量，实际是否无丢样仍需检查 STATUS DROP/MISS、RTT 换卡耗时与 meta.txt 的统计。",
            warning=True,
        )

        self.recipe(
            "功能 09：格式化指定或全部 SD 卡",
            "格式化单张物理卡槽，或在用户强确认后格式化全部已插入卡；文件系统固定为 exFAT。",
            "SD_FORMAT SLOT=<1..4> FS=EXFAT CONFIRM=YES\\n\n"
            "SD_FORMAT SLOT=ALL FS=EXFAT CONFIRM=ERASE_ALL\\n",
            "> SD_FORMAT SLOT=2 FS=EXFAT CONFIRM=YES\n"
            "< SEQ=<n> STATUS=0 MSG=SD_FORMAT_OK slot=2 fs=EXFAT\n"
            "  capacity_mb=<MB> free_mb=<MB>\n"
            "> SD_FORMAT SLOT=ALL FS=EXFAT CONFIRM=ERASE_ALL\n"
            "< SEQ=<n> STATUS=0 MSG=SD_FORMAT_ALL_OK formatted=2 skipped=2 failed=0 mask=0x05",
            "仅 IDLE 可用。先 SD_SCAN 再让用户核对卡槽和文件系统，弹出二次确认后才发 CONFIRM=YES。"
            "单卡未确认返回 SD_FORMAT_NEED_CONFIRM；全部格式化若未使用 ERASE_ALL 返回 SD_FORMAT_ALL_NEED_CONFIRM；忙时返回 SD_FORMAT_BUSY。"
            "全部格式化仅处理扫描到的已插入卡，空卡槽计入 skipped。GUI 必须使用单独的危险按钮、二次确认并等待最终响应，不能后台自动执行。",
        )

        self.recipe(
            "功能 10：对指定卡写入测试数据",
            "验证卡槽、文件系统和写入链路，不使用真实 ADC。",
            "SD_DEMO SLOT=<1..4|ALL> BYTES=<payload_bytes>\\n",
            "> SD_DEMO SLOT=2 BYTES=4096\n"
            "< SEQ=<n> STATUS=0 MSG=SD_DEMO tried=1 ok=1 fail=0 skip=0 bytes=4096",
            "仅 IDLE 可用。BYTES 指假数据负载，不含测试文件的文本头；固件上限为 65536 字节。"
            "测试文件在 /HYDROPHONE/TEST/，可用 5002 LIST/GET 或拔卡核对。"
            "ALL 会遍历选中卡，未就绪卡计入 skip；不能仅凭 STATUS=0 推断四张卡都通过。",
        )

        self.recipe(
            "功能 11：新增定时/重复采集任务",
            "授时后在设备 RAM 中注册按绝对时间写 SD 的采集任务；添加后自动等待执行，无 START_TASK 命令。",
            "TASK_ADD ID=<id> START=<unix_ms> DURATION=<ms> INTERVAL=<ms>\n"
            "         REPEAT=<count> RATE=<Hz> PREC=<bits> GAIN=<dB>\\n",
            "> TASK_ADD ID=1001 START=1789351500000 DURATION=60000\n"
            "  INTERVAL=300000 REPEAT=3 RATE=200000 PREC=24 GAIN=0\n"
            "< SEQ=<n> STATUS=0 MSG=TASK_ADD",
            "示例首次计划北京时间 10:05，采 60 秒，后两次计划开始于 10:10 和 10:15。"
            "发送时必须把 > 后两行合并为一条 ASCII 命令并仅在末尾加 LF；长度须小于控制口 128 字节。"
            "TASK_ADD 前需 TIME；无有效时间返回 TASK_ADD_TIME_INVALID。表最多 16 项，重启即丢失。",
        )
        self.note(
            "任务参数",
            "START 是 Unix 毫秒；DURATION/INTERVAL 均为毫秒，INTERVAL 是相邻计划开始间隔，不是前次结束后的等待。"
            "REPEAT=1 表示仅一次。固件未做任务冲突仲裁，也不禁止重复 ID；上位机必须预先检查时间窗不重叠、ID 不重复，"
            "并确认任务运行前 SD 卡可写。",
            warning=True,
        )

        self.recipe(
            "功能 12：读取已布置任务与当前执行状态",
            "设备重连后重建任务表界面，并判断是否正在执行。",
            "TASK_LIST\\n\nTASK_GET INDEX=<0..15>\\n",
            "> TASK_LIST\n< SEQ=<n> STATUS=0 MSG=TASK_LIST COUNT=1 MAX=16 TIME=1\n"
            "> TASK_GET INDEX=0\n"
            "< SEQ=<n> STATUS=0 MSG=TASK I=0 RUN=0 CMP=0 ID=1001\n"
            "  ST=1789351500000 END=1789351560000 DUR=60000 INT=300000\n"
            "  REP=3 EXE=0 RATE=200000 PREC=24 GAIN=0",
            "先 TASK_LIST，再对 INDEX=0..15 逐项 TASK_GET；空项返回 TASK_EMPTY I=<index>。"
            "不能只读取前 COUNT 项，因为删除后索引可能有空洞。I 是表索引，ID 是客户任务号；"
            "END 仅表示首轮计划结束，重复轮次需按 ST+INT*k 推算。",
        )

        self.recipe(
            "功能 13：删除单项或清空全部任务",
            "移除指定表索引，或清空 16 个表位。",
            "TASK_DELETE INDEX=<0..15>\\n\nTASK_CLEAR\\n",
            "> TASK_DELETE INDEX=0\n< SEQ=<n> STATUS=0 MSG=TASK_DELETE INDEX=0\n"
            "> TASK_CLEAR\n< SEQ=<n> STATUS=0 MSG=TASK_CLEAR",
            "删除使用 TASK_GET 返回的 I，不是 ID。删除空项或非法索引时 STATUS=-1。"
            "TASK_CLEAR 会停止正在运行的任务，属于高风险操作，GUI 应二次确认；"
            "完成后重新 TASK_LIST/TASK_GET 刷新界面。",
        )

        self.recipe(
            "功能 14：PC13 低电平触发后写 SD",
            "配置触发专用采样率、精度、增益和持续时间，布防后等待 PC13 稳定低电平 10 ms。",
            "TRIG_CONFIG OUTPUT=STORAGE REPEAT=<ONCE|CONTINUOUS> RATE=<Hz> PREC=<bits>\n"
            "            GAIN=<dB> DURATION_MS=<1..86400000>\\n\nTRIG_ARM\\n",
            "> TRIG_CONFIG OUTPUT=STORAGE REPEAT=CONTINUOUS RATE=96000 PREC=24\n"
            "  GAIN=0 DURATION_MS=10000\n"
            "< SEQ=<n> STATUS=0 MSG=TRIG_CONFIG_OK level=LOW\n"
            "  output=STORAGE repeat=CONTINUOUS rate=96000 prec=24BIT gain=+0DB duration_ms=10000\n"
            "> TRIG_ARM\n< SEQ=<n> STATUS=0 MSG=TRIG_ARM level=LOW repeat=CONTINUOUS",
            "两行 TRIG_CONFIG 只是排版，发送时合并成一条。仅 IDLE 可配置/布防。"
            "先 SD_SCAN 确认可写卡，再 TRIG_STATUS 确认 pin=1；若布防时 PC13 已低，也会在 10 ms 后触发。"
            "REPEAT=ONCE 在一次采集后回到 IDLE；REPEAT=CONTINUOUS 会等待 PC13 释放为高电平后自动重新布防。",
        )

        self.recipe(
            "功能 15：PC13 触发后实时传输及取消布防",
            "触发后把数据送到 5001，而不是写卡；可在等待或采集过程中取消。",
            "TRIG_CONFIG OUTPUT=REALTIME REPEAT=<ONCE|CONTINUOUS> RATE=<Hz> PREC=<bits>\n"
            "            GAIN=<dB> DURATION_MS=<ms>\\n\nTRIG_ARM\\n\nTRIG_STOP\\n",
            "> TRIG_CONFIG OUTPUT=REALTIME REPEAT=CONTINUOUS RATE=96000 PREC=24\n"
            "  GAIN=0 DURATION_MS=5000\n"
            "< SEQ=<n> STATUS=0 MSG=TRIG_CONFIG_OK level=LOW\n"
            "  output=REALTIME repeat=CONTINUOUS rate=96000 prec=24BIT gain=+0DB duration_ms=5000\n"
            "> TRIG_ARM\n< SEQ=<n> STATUS=0 MSG=TRIG_ARM level=LOW repeat=CONTINUOUS\n"
            "> TRIG_STOP\n< SEQ=<n> STATUS=0 MSG=TRIG_STOP",
            "先建立 5001 接收连接，再布防。TRIG_STATUS 可读 mode、pin、output、rate、prec、gain、duration_ms；"
            "触发后状态为 REALTIME_TRIGGERED，写卡触发则为 STORAGE_TRIGGERED。"
            "TRIG_STOP 可取消等待或提前停止，随后用 STATUS 确认 IDLE。",
        )

        self.recipe(
            "功能 16：通过网口列出并下载 SD 文件",
            "在 IDLE 下通过独立的 TCP 5002 浏览卡内目录并下载原始文件。",
            "LIST SLOT=<1..4> PATH=/HYDROPHONE\\n\n"
            "GET SLOT=<1..4> PATH=/HYDROPHONE/.../data.bin\\n",
            "> LIST SLOT=1 PATH=/HYDROPHONE\n"
            "< OK LIST SLOT=1 PATH=/HYDROPHONE\n"
            "< ENTRY TYPE=D NAME=DATA SIZE=0\n< END LIST COUNT=1\n"
            "> GET SLOT=1 PATH=/HYDROPHONE/DATA/<run>/data.bin\n"
            "< OK FILE SLOT=1 PATH=... SIZE=<N>\n< <紧接 N 字节原始文件>",
            "此处的 LIST 和 GET 发到 5002，不走 SEQ/STATUS/MSG。目录名应先 LIST 逐级获取，"
            "<run> 只是占位符，不能原样发送。GET 的首行结束后立即按 SIZE 接收 N 字节；"
            "二进制中可能含 LF，绝不能继续用 readline 读文件。下载完对比 SIZE，异常中断删除或标记残缺文件。",
        )

        self.heading("从零实现的收发骨架", 2)
        self.paragraph("控制端的每条请求采用一发一收；TCP 无消息边界，下面两个接收函数展示最小的安全处理方式。")
        self.code(
            "def recv_exact(sock, count):\n"
            "    buf = bytearray()\n"
            "    while len(buf) < count:\n"
            "        chunk = sock.recv(count - len(buf))\n"
            "        if not chunk:\n"
            "            raise ConnectionError('socket closed')\n"
            "        buf.extend(chunk)\n"
            "    return bytes(buf)\n\n"
            "def recv_line(sock, pending):\n"
            "    while b'\\n' not in pending:\n"
            "        chunk = sock.recv(4096)\n"
            "        if not chunk:\n"
            "            raise ConnectionError('socket closed')\n"
            "        pending.extend(chunk)\n"
            "    end = pending.index(b'\\n') + 1\n"
            "    line = bytes(pending[:end]).rstrip(b'\\r\\n')\n"
            "    del pending[:end]\n"
            "    return line"
        )
        self.bullets([
            "控制连接维护一个独立的 pending 缓冲区；每次 sendall(command.encode('ascii') + b'\\n') 后调用 recv_line，检查 STATUS 是否为 0。",
            "实时数据连接只使用 recv_exact 读取 16 字节头和 payload_len 字节负载，不能与控制 pending 混用。",
            "文件连接使用另一个 pending：先 recv_line 读 OK FILE，再先消耗 pending 中已有的二进制字节，最后读满 SIZE；或始终使用同一个 socket.makefile('rb')。",
            "GUI 不应在主线程做 socket.recv、卡扫描或大文件 GET；数据线程先持续读取和校验，再以批量方式更新波形。",
        ])
        self.paragraph("控制调用示例：同一个 pending 在连接的整个生命周期内复用，不能每条命令重新创建。")
        self.code(
            "def control_call(sock, pending, command):\n"
            "    sock.sendall(command.encode('ascii') + b'\\n')\n"
            "    line = recv_line(sock, pending).decode('ascii')\n"
            "    prefix, message = line.split(' MSG=', 1)\n"
            "    fields = dict(part.split('=', 1) for part in prefix.split())\n"
            "    seq = int(fields['SEQ'])\n"
            "    status = int(fields['STATUS'])\n"
            "    if status != 0:\n"
            "        raise RuntimeError(f'{seq}: {message}')\n"
            "    return seq, message"
        )
        self.paragraph("文件下载示例：按 SIZE 流式写入，不把 2 TiB 文件一次读进内存。")
        self.code(
            "import socket\n\n"
            "def download(ip, slot, remote, local):\n"
            "    with socket.create_connection((ip, 5002), timeout=5) as sock:\n"
            "        sock.settimeout(20)\n"
            "        sock.sendall(f'GET SLOT={slot} PATH={remote}\\n'.encode('ascii'))\n"
            "        with sock.makefile('rb') as stream:\n"
            "            head = stream.readline().decode('ascii').strip()\n"
            "            if not head.startswith('OK FILE '):\n"
            "                raise RuntimeError(head)\n"
            "            left = int(head.rsplit(' SIZE=', 1)[1])\n"
            "            with open(local, 'wb') as output:\n"
            "                while left:\n"
            "                    chunk = stream.read(min(left, 65536))\n"
            "                    if not chunk:\n"
            "                        raise EOFError(f'missing {left} bytes')\n"
            "                    output.write(chunk)\n"
            "                    left -= len(chunk)"
        )
        self.note(
            "异常处理",
            "控制命令超时可能是解析丢弃、网络断开，也可能是操作尚未完成；不要无条件重发格式化、任务添加或 START。"
            "文件下载异常时将本地文件标记为不完整。TCP 5000、5001、5002 的缓冲区和生命周期必须独立管理。",
            warning=True,
        )

    def section_overview(self):
        self.heading("1 快速对接总览", 1)
        self.note(
            "最短联调路径",
            "电脑设置为 192.168.1.x/24（不要占用 .80）→ 连接 5001 数据端口和 5000 控制端口 → 发送 STATUS → 发送 CONFIG → 发送 START_RT → 按实时包格式解析 → 发送 STOP_RT。",
        )

        self.heading("1.1 默认网络参数", 2)
        self.table(
            ["项目", "当前值", "上位机要求"],
            [
                ("设备 IPv4", "192.168.1.80", "电脑需处于 192.168.1.0/24 且地址不冲突"),
                ("子网掩码", "255.255.255.0", "直接连接或交换机连接均可"),
                ("网关", "192.168.1.1", "同网段直连时不依赖网关"),
                ("MAC", "02:20:0A:20:00:01", "多设备部署时必须保证 MAC/IP 唯一"),
                ("控制端口", "TCP 5000", "ASCII 文本命令和文本响应"),
                ("实时数据端口", "TCP 5001", "连续二进制采样包"),
                ("文件端口", "TCP 5002", "SD 目录 LIST 和文件 GET"),
            ],
            [1800, 2340, 5220],
            code_columns=(1,),
        )

        self.heading("1.2 功能模块与接口", 2)
        self.table(
            ["功能模块", "入口", "主要输出", "数据去向"],
            [
                ("设备状态与授时", "STATUS / TIME", "模式、连接、卡状态、授时结果", "控制端口"),
                ("采集配置", "CONFIG", "目标采样率、精度、增益配置结果", "控制端口"),
                ("实时采集", "START_RT / STOP_RT", "HY 二进制采样包", "5001 上位机"),
                ("立即自存储", "START_STORE / STOP_STORE", "meta.txt + data.bin", "SD1→SD4"),
                ("任务表采集", "TASK_*", "按绝对时间自动写卡", "SD1→SD4"),
                ("低电平触发", "TRIG_*", "实时上传或写卡", "5001 或 SD"),
                ("SD 管理", "SD_STATUS / SD_SCAN / SD_FORMAT / SD_DEMO", "卡槽、容量、格式化及测试结果", "控制端口"),
                ("文件读取", "LIST / GET", "目录文本或原始文件字节", "5002 上位机"),
            ],
            [1800, 2250, 3060, 2250],
            code_columns=(1,),
            small=True,
        )

        self.heading("1.3 四种采集方式", 2)
        self.table(
            ["方式", "启动条件", "停止条件", "是否写 SD"],
            [
                ("实时采集", "数据端口已连接，START_RT", "STOP_RT", "否"),
                ("立即自存储", "至少一张可写卡，START_STORE", "STOP_STORE 或写卡失败", "是"),
                ("任务表", "已授时且已添加启用任务", "任务持续时间到达", "是"),
                ("低电平触发", "TRIG_CONFIG 后 TRIG_ARM，PC13 低电平稳定 10 ms", "设定时长到达或 TRIG_STOP", "可选"),
            ],
            [1620, 3600, 2520, 1620],
        )

    def section_transport(self):
        self.heading("2 TCP 连接与通用协议", 1)
        self.heading("2.1 连接规则", 2)
        self.bullets([
            "三个端口是相互独立的 TCP Server。当前 W5500 配置为每个功能一个 socket，建议每个端口只连接一个上位机客户端。",
            "控制端口和文件端口采用按行文本；命令区分大小写，规范命令全部使用大写 ASCII。",
            "每条文本命令必须以 LF（\\n）或 CRLF（\\r\\n）结束。建议一发一收，收到响应后再发送下一条命令。",
            "控制接收缓冲区为 128 字节。为保留行结束符，单条命令总长度建议不超过 127 字节。",
            "TCP 是字节流。实时端口一次 recv 可能只得到半包，也可能包含多个包，上位机必须使用持续缓冲区拆包。",
            "当前协议没有身份认证、加密和远程访问控制，应部署在受控局域网内。",
        ])

        self.heading("2.2 控制响应格式", 2)
        self.code("SEQ=<设备序号> STATUS=<状态码> MSG=<消息正文>\\r\\n")
        self.code("SEQ=1 STATUS=0 MSG=IDLE CTRL=1 DATA=1 SDW=2 REC=0 SLOT=0 RUN=0 SAMPLES=0 TRIG_PIN=1 TOUT=STORAGE TDUR=10000 BUF=0 DROP=0 MISS=0 SW=0")
        self.table(
            ["字段", "含义"],
            [
                ("SEQ", "由设备对成功解析的控制命令递增分配；上位机不在请求中携带序号。设备重启后从 1 开始。"),
                ("STATUS=0", "命令处理成功。"),
                ("STATUS=-1", "命令已识别，但前置条件、参数或底层操作失败。具体信息见 MSG。"),
                ("MSG", "命令相关的结果正文，可能包含空格分隔的 KEY=VALUE 字段。"),
            ],
            [1800, 7560],
            code_columns=(0,),
        )
        self.paragraph("解析响应时先匹配固定前缀 SEQ=<整数> STATUS=<整数> MSG=，MSG 之后整段均属于正文；不要简单按空格只取 MSG 的第一个单词。")
        self.note(
            "无响应判定",
            "未知命令、非法参数或超长命令可能在解析阶段被丢弃，当前固件不会统一返回 BAD_COMMAND。超时后应检查大小写、换行符、字段名和总长度，不要无限等待。",
            warning=True,
        )

        self.heading("2.3 设备模式字符串", 2)
        self.table(
            ["模式字符串", "含义"],
            [
                ("IDLE", "空闲，可进行格式化、Demo 和文件读取。"),
                ("REALTIME_STREAM", "普通实时采集。"),
                ("STORAGE_IMMEDIATE", "立即自存储。"),
                ("STORAGE_SCHEDULED", "任务表触发的写卡采集。"),
                ("TRIGGER_ARMED", "低电平触发已布防，等待 PC13 变低。"),
                ("REALTIME_TRIGGERED", "触发后的实时上传正在运行。"),
                ("STORAGE_TRIGGERED", "触发后的 SD 写卡正在运行。"),
                ("FILE_SERVICE", "已预留状态；当前文件服务实际只在 IDLE 中同步执行。"),
                ("ERROR", "存储写入失败且无可切换卡等不可恢复错误。"),
            ],
            [2700, 6660],
            code_columns=(0,),
        )

        self.heading("2.4 最小控制客户端示例", 2)
        self.code(
            "import socket\n\n"
            "with socket.create_connection((\"192.168.1.80\", 5000), timeout=3) as s:\n"
            "    s.sendall(b\"STATUS\\n\")\n"
            "    response = b\"\"\n"
            "    while b\"\\n\" not in response:\n"
            "        response += s.recv(1024)\n"
            "    print(response.split(b\"\\n\", 1)[0].decode(\"ascii\", errors=\"replace\"))"
        )

    def section_config(self):
        self.heading("3 采集参数与增益配置", 1)
        self.heading("3.1 CONFIG 命令", 2)
        self.code("CONFIG RATE=200000 PREC=24 GAIN=0")
        self.table(
            ["参数", "允许值", "说明"],
            [
                ("RATE", "1/2/4/8/12/16/24/48/96/128/200 kS/s", "目标采样率。实时上传仅允许到 96 kS/s；ADC 时钟整数分频会造成少量实际误差。"),
                ("PREC", "16 / 24", "本版上位机对外支持的传输或存储样本位宽；不要在产品界面提供 32 bit。"),
                ("GAIN", "-30 / -20 / -10 / 0 / 20 / 40", "总增益挡位。-10/-30 dB 为数字衰减，其余为 PB6～PB9 硬件支路。"),
            ],
            [1500, 2340, 5520],
            code_columns=(0, 1),
        )
        self.paragraph("成功响应示例：")
        self.code("SEQ=2 STATUS=0 MSG=CONFIG_OK rate=200000 prec=24BIT gain=+0DB hw=+0DB sw=+0DB branch=0x01")
        self.paragraph("参数无法通过校验时：")
        self.code("SEQ=2 STATUS=-1 MSG=CONFIG_INVALID")

        self.heading("3.2 标称增益映射", 2)
        self.table(
            ["GAIN", "实现方式", "软件系数", "支路", "PB9..PB6 电平", "硬件增益"],
            [
                ("-30 dB", "软件衰减", "0.0316228", "PB6", "1110", "0 dB"),
                ("-20 dB", "硬件支路", "1.0", "PB9", "0111", "-20 dB"),
                ("-10 dB", "软件衰减", "0.3162278", "PB6", "1110", "0 dB"),
                ("0 dB", "硬件支路", "1.0", "PB6", "1110", "0 dB"),
                ("+20 dB", "硬件支路", "1.0", "PB7", "1101", "+20 dB"),
                ("+40 dB", "硬件支路", "1.0", "PB8", "1011", "+40 dB"),
            ],
            [1260, 1560, 1260, 1080, 1980, 2220],
            code_columns=(1, 3, 4),
            small=True,
        )
        self.note(
            "默认值",
            "上电默认 RATE=200000、PREC=24、GAIN=0 dB，对应 PB6 选中、PB9..PB6=1110。PB6～PB9 的支路控制为低电平有效，PC4/PC5 保持未配置状态。",
        )
        self.note(
            "配置生效时机",
            "CONFIG 只更新普通实时采集和立即自存储的“下一次启动配置”。正在运行的采集不会被在线重配置。任务表和触发模式在各自命令中携带独立配置。",
            warning=True,
        )

        self.heading("3.3 兼容写法", 2)
        self.paragraph("解析器还接受 SET_ACQ、SR、BITS 等别名和位置参数，但第三方上位机应固定使用规范 KEY=VALUE 写法，便于版本管理和日志定位。ATT、RANGE、SW 已不再接受。")
        self.code("SET_ACQ RATE=200000 BITS=24 GAIN=0    # 可解析，但不建议作为主协议写法")

    def section_status_time(self):
        self.heading("4 状态查询与上位机授时", 1)
        self.heading("4.1 STATUS", 2)
        self.code("STATUS")
        self.table(
            ["返回字段", "说明"],
            [
                ("首个单词", "当前系统模式，例如 IDLE、REALTIME_STREAM。"),
                ("CTRL", "控制端口是否已建立连接，1=已连接。"),
                ("DATA", "实时数据端口是否已建立连接，1=已连接。"),
                ("SDW", "最近一次 SD 扫描得到的可写卡数量。"),
                ("REC", "存储记录器是否正在写卡。"),
                ("SLOT", "记录器当前或最近使用的卡槽；REC=0 时不能据此判断正在写卡。"),
                ("RUN", "本次上电期间分配的记录运行号。"),
                ("SAMPLES", "记录器累计样本数的低 32 位，仅在写卡模式下有意义。"),
                ("TRIG_PIN", "PC13 当前电平：0=低，1=高。"),
                ("TOUT", "触发后动作：STORAGE 或 REALTIME。"),
                ("TDUR", "触发采集持续时间，单位 ms。"),
                ("BUF", "当前采样环形缓冲中的待处理样本数；容量为 163839 个样本。"),
                ("DROP", "本次采集因环形缓冲写满而丢弃的样本数低 32 位。"),
                ("MISS", "根据 DRDY 周期间隔估算的漏采样本数低 32 位。"),
                ("SW", "本次存储运行已经完成的跨卡切换次数。"),
            ],
            [1980, 7380],
            code_columns=(0,),
        )

        self.heading("4.2 TIME 授时", 2)
        self.code("TIME 1789351200000")
        self.paragraph("参数为 Unix Epoch 毫秒。设备记录接收时刻的 HAL tick，之后用“授时值 + 已运行毫秒”生成当前时间。成功响应：")
        self.code("SEQ=3 STATUS=0 MSG=TIME_OK")
        self.paragraph("北京时间 2026-09-14 10:00:00 的显式时区转换示例：")
        self.code(
            "from datetime import datetime, timezone, timedelta\n"
            "cst = timezone(timedelta(hours=8))\n"
            "dt = datetime(2026, 9, 14, 10, 0, 0, tzinfo=cst)\n"
            "unix_ms = int(dt.timestamp() * 1000)"
        )
        self.bullets([
            "时间戳本身不含时区。上位机应先确认电脑时区和系统时间正确，再做转换。",
            "设备复位或掉电后授时失效；任务表也会被清空，必须重新授时和下发任务。",
            "立即自存储不会强制要求时间有效。若未授时，meta.txt 中的时间是上电 tick，不是可还原的日历时间。",
            "任务添加要求设备已经授时，否则返回 TASK_ADD_TIME_INVALID。",
        ])

    def section_realtime(self):
        self.heading("5 实时采集与二进制数据流", 1)
        self.heading("5.1 推荐操作顺序", 2)
        self.numbered([
            "先连接 TCP 5001 数据端口，并启动持续接收线程。",
            "连接 TCP 5000 控制端口，发送 STATUS，确认 DATA=1。",
            "发送 CONFIG，并确认 STATUS=0。",
            "发送 START_RT；成功消息为 START_REALTIME。",
            "按 HY 包头持续拆包、解码、统计序号和速率。",
            "结束时发送 STOP_RT；成功消息为 STOP_REALTIME。",
        ])
        self.code("START_RT\n# 成功：SEQ=<n> STATUS=0 MSG=START_REALTIME\n# 失败：SEQ=<n> STATUS=-1 MSG=START_REALTIME_FAIL")
        self.note(
            "数据连接是硬前置条件",
            "若 5001 未连接，START_RT 会失败。实时过程中若 5001 断开，设备会自动停止采集；若上位机长期不读并造成首次缓冲溢出，设备也会停止实时模式。上位机应监控连接和 STATUS，并在重连后重新启动。",
            warning=True,
        )

        self.heading("5.2 实时包头", 2)
        self.caption("表 5-1  实时数据包固定 16 字节包头（全部多字节字段为小端）")
        self.table(
            ["偏移", "长度", "字段", "类型", "说明"],
            [
                ("0", "2", "magic", "char[2]", "固定为 ASCII 'H','Y'。"),
                ("2", "1", "version", "uint8", "当前为 1。"),
                ("3", "1", "precision", "uint8", "本版上位机接受 16 或 24。"),
                ("4", "4", "sequence", "uint32 LE", "设备本次上电期间单调递增，uint32 回绕；连续触发分段不清零。"),
                ("8", "4", "sample_rate", "uint32 LE", "定时器量化后的实际采样率，不是单纯回显目标值。"),
                ("12", "2", "sample_count", "uint16 LE", "本包样本数，当前最大 256。"),
                ("14", "2", "payload_len", "uint16 LE", "负载字节数，应等于 sample_count × 每样本字节数。"),
                ("16", "payload_len", "payload", "signed LE", "连续有符号样本，无时间戳、无应用层 CRC。"),
            ],
            [900, 900, 1440, 1620, 4500],
            code_columns=(0, 1, 2, 3),
            small=True,
        )
        self.code('Python struct：PACKET_HEADER = struct.Struct("<2sBBIIHH")')

        self.heading("5.3 样本负载", 2)
        self.table(
            ["PREC", "每样本", "码值范围", "解码"],
            [
                ("16", "2 字节", "-32768～32767", "int16，小端"),
                ("24", "3 字节", "-8388608～8388607", "24 位二进制补码，小端，需手动符号扩展"),
            ],
            [1080, 1440, 2700, 4140],
            code_columns=(0,),
        )
        self.code(
            "def decode_samples(payload, precision, count):\n"
            "    width = {16: 2, 24: 3}[precision]\n"
            "    if len(payload) != count * width:\n"
            "        raise ValueError(\"payload length mismatch\")\n"
            "    out = []\n"
            "    for i in range(count):\n"
            "        p = i * width\n"
            "        if precision == 16:\n"
            "            out.append(int.from_bytes(payload[p:p+2], \"little\", signed=True))\n"
            "        elif precision == 24:\n"
            "            raw = int.from_bytes(payload[p:p+3], \"little\", signed=False)\n"
            "            out.append(raw - 0x1000000 if raw & 0x800000 else raw)\n"
            "    return out"
        )

        self.heading("5.4 电压换算", 2)
        self.paragraph("当前验证 GUI 按 VREF=4.096 V、输入范围系数 1.0，把 ADC 差分码值换算为：")
        self.code("16 bit: Vdiff = code / 32768   × 4.096 V\n24 bit: Vdiff = code / 8388608 × 4.096 V\nVpp = max(Vdiff) - min(Vdiff)")
        self.note(
            "测量口径",
            "这里表示 ADC 差分输入 AN1+ - AN1- 的码值换算。实际整机输入电压还需要结合前端标称增益、元件误差、参考源误差和系统标定系数；不要把 GUI 的 V 值直接当作水听器声压值。",
            warning=True,
        )

        self.heading("5.5 理论应用层数据量", 2)
        self.table(
            ["采样率", "精度", "纯样本负载", "按 256 点/包含包头", "约合比特率"],
            [
                ("96 kS/s 实时", "16 bit", "187.5 KiB/s", "193.4 KiB/s", "1.58 Mbit/s"),
                ("96 kS/s 实时", "24 bit", "281.3 KiB/s", "287.1 KiB/s", "2.35 Mbit/s"),
                ("200 kS/s 写卡", "16 bit", "390.6 KiB/s", "不适用", "3.20 Mbit/s"),
                ("200 kS/s 写卡", "24 bit", "585.9 KiB/s", "不适用", "4.80 Mbit/s"),
            ],
            [1500, 1260, 2100, 2700, 1800],
            small=True,
        )
        self.paragraph("上表不含 Ethernet/IP/TCP 帧开销，且假设每包恰好 256 点。实际吞吐应以上位机每秒收到的样本数和设备日志中的 measured/tx rate 为准。")

        self.heading("5.6 上位机拆包要点", 2)
        self.bullets([
            "在累计缓冲区中搜索 HY；不足 16 字节时继续接收。",
            "解析包头后验证 version、precision、sample_count 和 payload_len，未收齐整包时继续等待。",
            "一个 recv 中可能包含多包，应循环解析到缓冲区不足一包为止。",
            "用 uint32 模运算监控 sequence 是否连续。序号只在设备初始化时置 0；普通停止/重启实时流和连续触发分段均不清零。",
            "包内没有设备时间戳。现有 GUI 保存 CSV 时写入的是电脑收包时刻，不能等同于每个样本的精确采样时刻。",
        ])

    def section_manual_storage(self):
        self.heading("6 立即自存储模式", 1)
        self.heading("6.1 推荐流程", 2)
        self.numbered([
            "插卡或换卡后发送 SD_SCAN，确认至少一个卡槽为 R。",
            "发送 TIME，确保文件元数据中的 start_time_ms 是有效 Unix 毫秒。",
            "发送 CONFIG，并确认参数。",
            "发送 START_STORE；成功消息为 START_STORAGE。",
            "采集期间不要调用 SD_SCAN、格式化、Demo 或文件服务。",
            "结束时发送 STOP_STORE；等待 STOP_STORAGE 响应后再拔卡或读取文件。",
        ])
        self.code("START_STORE\n# 成功：SEQ=<n> STATUS=0 MSG=START_STORAGE\nSTOP_STORE\n# 成功：SEQ=<n> STATUS=0 MSG=STOP_STORAGE")
        self.paragraph("启动失败时，MSG 返回当前 SD 摘要，而不是固定的 START_STORAGE_FAIL。例如：")
        self.code("SEQ=8 STATUS=-1 MSG=SD scan=1 present=0 writable=0 total_mb=0 S1=- S2=- S3=- S4=-")

        self.heading("6.2 卡槽选择与跨卡", 2)
        self.bullets([
            "记录器按物理卡槽 1→2→3→4 选择第一张可写卡。",
            "记录启动前扫描并缓存可写卡队列；采集中不再重新扫描四个卡槽，换卡直接尝试更高编号的缓存卡。",
            "有后续卡时，当前卡估算剩余约 16 MiB 即提前切换；正常换卡写入段尾元数据并同步关闭。",
            "写入或同步失败时快速放弃故障卡，不重复等待同步；已按完整样本写入的部分不会在下一张卡重复写。",
            "一次记录跨卡时，各卡上创建同名运行目录；meta.txt 中的 segment、slot、continued 用于合并。",
            "若当前卡之后没有可写卡，系统进入 ERROR。中途不会回绕到更低编号卡槽。",
            "记录过程中只给当前选中的卡保持供电，其他未使用卡处于断电状态。",
        ])

        self.heading("6.3 目录和文件", 2)
        self.code(
            "/HYDROPHONE/DATA/RUN_YYYYMMDD_HHMMSS_mmm_<RATE>Hz_MANUAL_<RUN_ID>/\n"
            "    meta.txt\n"
            "    data.bin"
        )
        self.paragraph("目录名直接包含本地开始时间、配置采样率、采集来源和运行编号。设备未授时时使用 RUN_UNSYNC_<TICK_HEX>_<RATE>Hz_<SOURCE>_<RUN_ID>，准确时间和其他采样设置仍以 meta.txt 为准。")

    def section_tasks(self):
        self.heading("7 任务表模式", 1)
        self.note(
            "当前启动逻辑",
            "任务表没有单独的 START_TASK 命令。设备授时后，TASK_ADD 成功即表示该启用任务进入等待状态；到达开始时间会自动采集并写 SD。",
        )

        self.heading("7.1 添加任务", 2)
        self.code("TASK_ADD ID=1001 START=1789351500000 DURATION=60000 INTERVAL=300000 REPEAT=3 RATE=200000 PREC=24 GAIN=0")
        self.table(
            ["参数", "含义与规则"],
            [
                ("ID", "非 0 的 uint32 业务任务号。固件不检查 ID 是否重复。"),
                ("START", "首次开始时间，Unix 毫秒。若已在过去，调度循环会尽快启动。"),
                ("DURATION", "每次采集持续时间，毫秒，必须大于 0。"),
                ("INTERVAL", "相邻两次计划开始时间的间隔，毫秒；不是“上次结束后的等待时间”。"),
                ("REPEAT", "执行次数。现有 GUI 强制至少 1；原始协议中 0 会按无限重复处理。"),
                ("ENABLE", "可选，1=启用。解析默认值为 1；设为 0 会导致 TASK_ADD 校验失败。"),
                ("RATE/PREC/GAIN", "该任务独立的采样配置，允许值与 CONFIG 相同。"),
            ],
            [1800, 7560],
            code_columns=(0,),
        )
        self.code("# 未授时\nSEQ=<n> STATUS=-1 MSG=TASK_ADD_TIME_INVALID\n# 已授时且有空表项\nSEQ=<n> STATUS=0 MSG=TASK_ADD")

        self.heading("7.2 查询、删除和清空", 2)
        self.table(
            ["命令", "典型响应", "说明"],
            [
                ("TASK_LIST", "TASK_LIST COUNT=2 MAX=16 TIME=1", "返回数量、最大表项数和授时有效标志。"),
                ("TASK_GET INDEX=0", "TASK I=0 RUN=0 CMP=0 ...", "读取表项 0～15 的完整信息。"),
                ("TASK_GET INDEX=1", "TASK_EMPTY I=1", "该索引有效但为空。"),
                ("TASK_DELETE INDEX=0", "TASK_DELETE INDEX=0", "STATUS 表示是否成功；运行中的任务会先停止。"),
                ("TASK_CLEAR", "TASK_CLEAR", "停止所有运行任务并清空 16 个表项。"),
            ],
            [2160, 3960, 3240],
            code_columns=(0, 1),
            small=True,
        )
        self.paragraph("推荐读取方法：先 TASK_LIST，再依次发送 TASK_GET INDEX=0～15；根据 TASK、TASK_EMPTY 更新上位机表格。")

        self.heading("7.3 TASK 返回字段", 2)
        self.table(
            ["字段", "说明"],
            [
                ("I", "任务表索引 0～15；删除操作使用该索引，不是 ID。"),
                ("RUN", "1=本次任务正在采集。"),
                ("CMP", "1=已完成所有重复次数。"),
                ("ID", "上位机提供的任务 ID。"),
                ("ST / END", "首次开始和首次结束 Unix 毫秒；END=ST+DUR。"),
                ("DUR / INT", "单次持续时间和相邻计划开始间隔，毫秒。"),
                ("REP / EXE", "计划次数和已经完成的次数。"),
                ("RATE / PREC / GAIN", "该任务的采集参数，GAIN 为六个预定义硬件或软件挡位之一。"),
            ],
            [1800, 7560],
            code_columns=(0,),
        )

        self.heading("7.4 调度约束", 2)
        self.note(
            "避免时间重叠",
            "当前调度器没有任务排队和优先级仲裁。多个任务重叠，或任务到点时设备正处于普通实时/立即自存储模式，后启动的任务可能停止并替换当前采集。上位机下发前必须检查任务时间窗不重叠，并避免在任务窗口内手动启动其他模式。",
            warning=True,
        )
        self.bullets([
            "同一任务按 START + INTERVAL × 已执行次数计算下一次计划开始。",
            "若 INTERVAL 小于 DURATION，下一次在上一轮停止后立即启动，不会并行采集。",
            "任务在 TRIGGER_ARMED、REALTIME_TRIGGERED、STORAGE_TRIGGERED 期间暂停调度；退出触发状态后，已过期任务可能立即启动。",
            "任务表和已执行计数仅在 RAM 中，掉电不保存。",
            "任务写卡目录包含 _TASK_ 来源字段，meta.txt 中 source=TASK_TABLE、task_id=<ID>。",
        ])

    def section_trigger(self):
        self.heading("8 PC13 低电平触发模式", 1)
        self.heading("8.1 硬件和判定规则", 2)
        self.paragraph("PC13 由板上 4.7 kΩ 上拉到 3.3 V，外部光耦和三极管动作后把 PC13 拉低。当前固件只支持低电平触发，低电平连续保持至少 10 ms 才被接受。")
        self.note(
            "布防时已为低电平",
            "如果发送 TRIG_ARM 时 PC13 已经为低，设备会立即开始 10 ms 消抖并触发。需要“新边沿”语义时，上位机应先用 TRIG_STATUS 确认 pin=1，再布防。",
            warning=True,
        )

        self.heading("8.2 配置与布防", 2)
        self.code("TRIG_CONFIG OUTPUT=STORAGE REPEAT=CONTINUOUS RATE=200000 PREC=24 GAIN=0 DURATION_MS=10000\nTRIG_ARM")
        self.table(
            ["参数", "允许值", "说明"],
            [
                ("OUTPUT", "STORAGE / REALTIME", "触发后写 SD 或通过 5001 实时上传。兼容 STORE/SD、RT/TCP 等别名。"),
                ("REPEAT", "ONCE / CONTINUOUS", "单次采集后退出，或释放后自动重新布防。省略时为 ONCE。"),
                ("RATE", "1/2/4/8/12/16/24/48/96/128/200 kS/s", "触发专用目标采样率；OUTPUT=REALTIME 时最高 96 kS/s。"),
                ("PREC", "16 / 24", "触发专用精度；第三方上位机不开放 32 bit。"),
                ("GAIN", "-30/-20/-10/0/20/40", "触发专用总增益；软件衰减同样作用于实时和存储数据。"),
                ("DURATION_MS", "1～86400000", "触发后采集持续时间，最大 24 小时。"),
            ],
            [1800, 2520, 5040],
            code_columns=(0, 1),
        )
        self.bullets([
            "TRIG_CONFIG 仅在 IDLE 可用；否则返回 TRIG_CONFIG_BUSY。",
            "OUTPUT=STORAGE 布防前要求缓存的 writable_count 大于 0；换卡后应先 SD_SCAN。",
            "OUTPUT=REALTIME 布防前要求 5001 已连接。",
            "TRIG_ARM 成功后模式为 TRIGGER_ARMED。CONTINUOUS 在每段结束后先等待 PC13 恢复高电平，再接受下一次低电平。",
            "采集正在进行时到来的新触发不排队，PC13 持续低电平也不会重复启动。",
            "TRIG_STOP 可取消等待，也可提前停止触发后的采集。",
        ])

        self.heading("8.3 状态查询", 2)
        self.code("TRIG_STATUS")
        self.code("SEQ=<n> STATUS=0 MSG=TRIG_STATUS mode=TRIGGER_ARMED level=LOW pin=1 output=STORAGE repeat=CONTINUOUS count=3 wait=0 rate=200000 prec=24BIT gain=+0DB duration_ms=10000")
        self.table(
            ["字段", "说明"],
            [
                ("mode", "当前系统模式。"),
                ("level", "固定 LOW，表示只支持低电平触发。"),
                ("pin", "PC13 当前电平，0=低、1=高。"),
                ("output", "STORAGE 或 REALTIME。"),
                ("repeat", "ONCE 或 CONTINUOUS。"),
                ("count", "本次布防后已启动的触发采集次数。"),
                ("wait", "1=连续模式正等待 PC13 释放为高电平。"),
                ("rate / prec / gain", "最近一次触发配置。"),
                ("duration_ms", "触发后持续时间。"),
            ],
            [1980, 7380],
            code_columns=(0,),
        )
        self.paragraph("触发写卡目录包含 _TRIG_ 来源字段，meta.txt 中 source=TRIGGER。触发实时上传使用与普通实时采集完全相同的 HY 数据包。")

    def section_sd(self):
        self.heading("9 SD 卡管理", 1)
        self.heading("9.1 卡槽状态与扫描", 2)
        self.code("SD_STATUS    # 读取最近一次扫描缓存，不重新访问卡\nSD_SCAN      # 逐槽初始化、读取容量和剩余空间，并更新缓存")
        self.code("SEQ=<n> STATUS=0 MSG=SD scan=1 present=2 writable=1 total_mb=1953520 S1=R:1024000MB S2=U S3=- S4=-")
        self.table(
            ["卡槽标记", "含义"],
            [
                ("S1=R:xxxxMB", "卡已初始化、文件系统可挂载且可用于写入；数值为剩余容量 MB。"),
                ("S1=U", "检测到有效容量，但文件系统不可用、未格式化或挂载失败。"),
                ("S1=-", "未插卡，或卡初始化/容量读取失败。当前协议不能进一步区分这几种原因。"),
                ("S1=?", "尚未完成任何扫描。正常开机初始化会自动扫描，因此很少出现。"),
            ],
            [2520, 6840],
            code_columns=(0,),
        )
        self.bullets([
            "开机自动扫描 1～4 号槽。SD_SCAN 完成后所有卡槽断电。",
            "插拔卡不会自动刷新缓存；换卡后必须 SD_SCAN。",
            "present 是底层初始化和容量读取成功，不等同于可写；业务启动应检查 writable。",
            "代码容量模型为每卡最多 2 TiB、四卡总计最多 8 TiB；实际卡型和长期稳定性需实卡验证。",
        ])

        self.heading("9.2 格式化", 2)
        self.code("SD_FORMAT SLOT=1 FS=EXFAT CONFIRM=YES\nSD_FORMAT SLOT=ALL FS=EXFAT CONFIRM=ERASE_ALL")
        self.table(
            ["条件/字段", "规则"],
            [
                ("SLOT", "单卡为 1～4；全部已插入卡使用 ALL。"),
                ("FS", "固定为 EXFAT；固件不接受 FAT32 格式化请求。"),
                ("CONFIRM", "单卡必须 YES；全部卡必须使用更强确认 ERASE_ALL，YES 会返回 SD_FORMAT_ALL_NEED_CONFIRM。"),
                ("设备模式", "仅 IDLE 允许；其他模式返回 SD_FORMAT_BUSY。"),
                ("单卡成功", "SD_FORMAT_OK slot=<n> fs=EXFAT capacity_mb=<n> free_mb=<n>。"),
                ("全部成功", "SD_FORMAT_ALL_OK formatted=<n> skipped=<n> failed=0 mask=0x<xx>。"),
                ("部分失败", "STATUS=-1，MSG=SD_FORMAT_ALL_PARTIAL ... failed_mask=0x<xx>。"),
                ("没有卡", "STATUS=-1，MSG=SD_FORMAT_ALL_NO_CARD formatted=0 skipped=4 failed=0。"),
            ],
            [2160, 7200],
            code_columns=(0,),
        )
        self.note(
            "数据安全",
            "格式化会清空目标卡全部数据。单卡和全部卡必须使用不同按钮；全部卡确认框应明确写出“所有已插入卡都会被清空”，"
            "用户确认后发送 ERASE_ALL，并在等待期间禁用两个格式化按钮。收到最终响应后弹窗展示 formatted、skipped、failed；"
            "mask 的 bit0～bit3 分别代表卡槽 1～4。无论成功或失败都应重新 SD_SCAN 刷新界面。",
            warning=True,
        )

        self.heading("9.3 写入测试 Demo", 2)
        self.code("SD_DEMO SLOT=ALL BYTES=4096\nSD_DEMO SLOT=2 BYTES=65536\nSD_DEMO MASK=0x05 BYTES=4096")
        self.paragraph("BYTES=0 时使用 4096；超过 65536 会被固件截到 65536。未插卡或文件系统不可用的目标卡会被跳过。响应：")
        self.code("SEQ=<n> STATUS=0 MSG=SD_DEMO tried=2 ok=2 fail=0 skip=2 bytes=4096")
        self.paragraph("测试文件位于 /HYDROPHONE/TEST/SD<n>_DEMO_<TICK_HEX>.bin，包含文本头和递增字节模式。该命令只在 IDLE 可用，不会删除已有文件。")

        self.heading("9.4 供电策略", 2)
        self.table(
            ["操作", "卡槽供电行为"],
            [
                ("IDLE", "扫描完成后四个卡槽全部断电。"),
                ("SD_SCAN", "逐个选择并上电检测，结束后全部断电。"),
                ("格式化 / Demo / LIST / GET", "仅操作目标卡，操作结束后卸载并全部断电。"),
                ("自存储 / 任务写卡 / 触发写卡", "当前写入卡保持供电，其他卡断电；切卡时重新选择。"),
                ("实时上传", "不使用 SD，卡槽保持断电。"),
            ],
            [2700, 6660],
        )
        self.note(
            "并发禁令",
            "SD_SCAN、SD_FORMAT、SD_DEMO 和文件服务都只允许在 IDLE 执行。忙时 SD_SCAN 返回 SD_SCAN_BUSY，上位机应先停止当前模式并确认回到 IDLE。",
            warning=True,
        )

    def section_file_service(self):
        self.heading("10 SD 文件服务（TCP 5002）", 1)
        self.paragraph("文件服务与控制协议相互独立，不使用 SEQ/STATUS/MSG。建议每次 LIST 或 GET 新建一个 5002 连接，完成后关闭。文件服务只读，不支持上传、删除、重命名。")

        self.heading("10.1 LIST 目录读取", 2)
        self.code("LIST SLOT=1 PATH=/HYDROPHONE")
        self.code(
            "OK LIST SLOT=1 PATH=/HYDROPHONE\n"
            "ENTRY TYPE=D NAME=DATA SIZE=0\n"
            "ENTRY TYPE=D NAME=TEST SIZE=0\n"
            "END LIST COUNT=2"
        )
        self.table(
            ["字段", "说明"],
            [
                ("TYPE=D", "目录。"),
                ("TYPE=F", "普通文件。"),
                ("NAME", "当前目录下的名称。协议以空格分词，路径和文件名不要包含空格。"),
                ("SIZE", "文件字节数；目录通常为 0。"),
                ("COUNT", "本次返回的 ENTRY 数量。当前最多 64 项，无分页接口。"),
            ],
            [1800, 7560],
            code_columns=(0,),
        )

        self.heading("10.2 GET 文件下载", 2)
        self.code("GET SLOT=1 PATH=/HYDROPHONE/DATA/RUN_20260915_143020_123_200000Hz_MANUAL_000001/data.bin")
        self.code("OK FILE SLOT=1 PATH=/HYDROPHONE/.../data.bin SIZE=12345678\n<随后紧跟 12345678 字节原始文件内容>")
        self.numbered([
            "先按行读取响应头，不要把响应头后的二进制字节丢弃。",
            "从响应头解析 SIZE。",
            "把同一次 recv 中响应头之后的剩余字节作为文件数据。",
            "继续接收，直到累计字节数严格等于 SIZE。不要等待额外 END 标记。",
            "校验本地文件大小；网络中断或超时应删除/标记不完整文件。",
        ])
        self.code(
            "def recv_exact(sock, initial, size):\n"
            "    data = bytearray(initial[:size])\n"
            "    while len(data) < size:\n"
            "        chunk = sock.recv(min(8192, size - len(data)))\n"
            "        if not chunk:\n"
            "            raise ConnectionError(\"file transfer interrupted\")\n"
            "        data.extend(chunk)\n"
            "    return bytes(data)"
        )

        self.heading("10.3 文件服务错误", 2)
        self.table(
            ["返回", "常见原因"],
            [
                ("ERR BUSY", "系统不是 IDLE，例如正在实时、写卡、等待触发或触发采集。"),
                ("ERR BAD_LIST / ERR BAD_GET", "卡槽或路径语法非法，PATH 不是绝对路径，或包含 '..'。"),
                ("ERR MOUNT SLOT=n FR=x", "目标卡无法挂载。"),
                ("ERR LIST SLOT=n FR=x", "目录不存在或无法打开。"),
                ("ERR GET SLOT=n FR=x", "文件不存在或无法打开。"),
                ("ERR UNKNOWN", "不是 LIST 或 GET。"),
            ],
            [3420, 5940],
            code_columns=(0,),
        )
        self.bullets([
            "PATH 必须以 / 开头，禁止 '..'；当前生成目录不含空格。",
            "LIST 一次只返回前 64 项且没有继续页标记。目录文件超过 64 项时，应按日期/批次组织目录，或已知完整路径后直接 GET。",
            "GET 发送完 SIZE 指定的字节后没有结束行。客户端应以长度为准并主动关闭连接。",
            "GET 过程中发生 SD 读错误时，固件不会再发送结构化尾部错误；客户端必须依赖超时和长度校验。",
        ])

    def section_offline_data(self):
        self.heading("11 SD 记录文件的离线解析", 1)
        self.heading("11.1 运行目录来源", 2)
        self.table(
            ["目录前缀", "meta.txt source", "来源"],
            [
                ("目录含 _MANUAL_", "MANUAL", "START_STORE 立即自存储"),
                ("目录含 _TASK_", "TASK_TABLE", "任务表自动采集"),
                ("目录含 _TRIG_", "TRIGGER", "PC13 低电平触发后写 SD"),
            ],
            [2700, 2340, 4320],
            code_columns=(0, 1),
        )

        self.heading("11.2 meta.txt 字段", 2)
        self.table(
            ["字段", "说明"],
            [
                ("run_id", "本次上电期间递增的运行号，从 1 开始；跨重启可能重复。"),
                ("segment", "同一运行跨卡时的段号，从 1 开始。"),
                ("slot", "物理卡槽号 1～4。"),
                ("source", "MANUAL、TASK_TABLE 或 TRIGGER。"),
                ("task_id", "任务表来源时对应 ID；其他来源通常为 0。"),
                ("continued", "0=首段，1=从前一张卡续写。"),
                ("start_time_ms", "整个运行开始时间。已授时时为 Unix 毫秒，未授时时为上电 tick。"),
                ("sample_rate_hz / actual_sample_rate_hz", "请求的目标采样率和定时器量化后的实际采样率。离线时间轴应使用 actual_sample_rate_hz。"),
                ("precision / bytes_per_sample", "本版正常记录为 16BIT/2 或 24BIT/3。"),
                ("gain_db / hardware_gain_db / software_gain_db", "总标称增益、硬件支路增益和软件衰减；-10/-30 dB 会在公共采样缓冲前完成数字缩放。"),
                ("range_switch_bits", "硬件增益支路位，仅用于诊断；产品上位机只展示 GAIN，不提供支路位编辑。"),
                ("segment_first_sample_index", "本段第一个样本在整个运行中的零基索引，用于跨卡连续性校验。"),
                ("buffer_capacity_samples / card_free_at_open_bytes", "采样缓冲容量和打开本卡时的缓存剩余空间。"),
                ("file", "当前固定为 data.bin。"),
                ("segment_end_reason", "STOP、LOW_SPACE、WRITE_FAIL、SYNC_FAIL 或 CLOSE_FAIL。故障卡可能无法写入段尾。"),
                ("segment_stop_time_ms", "本段停止时间。"),
                ("segment_samples / segment_bytes / segment_trailing_bytes", "本段有效样本数、有效数据字节数及无法裁剪的残留非整样本字节。"),
                ("acq_*", "采集计数、环形缓冲丢样、估算漏采、SPI 读错和 ISR 最大周期/预算。"),
                ("previous_slot / previous_switch_*", "续写段记录上一卡槽、换卡原因和换卡耗时。"),
                ("run_stop_time_ms", "整个运行正常停止时间；通常只出现在最后一段。"),
                ("run_samples_written / run_bytes_written", "整个运行累计数量；通常只出现在最后一段。"),
            ],
            [2700, 6660],
            code_columns=(0,),
            small=True,
        )

        self.heading("11.3 data.bin", 2)
        self.paragraph("data.bin 没有包头和时间戳，是连续有符号小端样本。解码规则与实时 payload 完全一致。文件样本数应等于 file_size / bytes_per_sample。")
        self.code(
            "from pathlib import Path\n\n"
            "raw = Path(\"data.bin\").read_bytes()\n"
            "precision = 24\n"
            "width = {16: 2, 24: 3}[precision]\n"
            "assert len(raw) % width == 0\n"
            "samples = decode_samples(raw, precision, len(raw) // width)"
        )

        self.heading("11.4 跨卡合并", 2)
        self.numbered([
            "在四张卡中寻找同名 RUN_* 目录，并读取每份 meta.txt。",
            "校验 source、run_id、start_time_ms、sample_rate_hz、precision、gain_db 一致。",
            "按 segment 从小到大排序。",
            "逐段拼接 data.bin；每段大小应与 segment_bytes 对应。",
            "用最后一段的 run_samples_written/run_bytes_written 做总量校验。",
        ])
        self.note(
            "实际采样率与连续性",
            "实时包头和 meta.txt 都记录实际采样率。每秒 RTT 日志还给出 measured、DRDY、buffered、dropped、missed、errors 和 isr_max/budget；验收时应确认错误计数为 0 且 ISR 最大周期低于预算并留有余量。",
        )

    def section_concurrency(self):
        self.heading("12 模式互斥与上位机状态管理", 1)
        self.caption("表 12-1  当前实现下的命令并发规则")
        self.table(
            ["命令/功能", "可用条件", "对当前采集的影响"],
            [
                ("STATUS / SD_STATUS", "任意模式", "只读取内存状态，不停止采集。"),
                ("TIME", "任意模式", "更新软件时间基准，不停止采集。"),
                ("CONFIG", "任意模式可解析", "只影响下一次普通启动；不在线改变当前采集。"),
                ("START_RT", "5001 已连接", "先停止当前活动采集，再启动普通实时。"),
                ("START_STORE", "存在可写卡", "先停止当前活动采集，再启动立即自存储。"),
                ("STOP_RT / STOP_STORE", "任意模式", "当前实现调用统一停止函数，可能停止另一种活动采集。"),
                ("TASK_ADD", "已授时", "任务到点后可能停止并替换普通实时或立即自存储。"),
                ("TRIG_CONFIG / TRIG_ARM", "仅 IDLE", "布防期间暂停任务调度。"),
                ("TRIG_STOP", "任意模式可发送", "只在触发相关模式中停止触发采集。"),
                ("SD_FORMAT / SD_DEMO", "仅 IDLE", "忙时返回失败。"),
                ("LIST / GET", "仅 IDLE", "忙时返回 ERR BUSY。"),
                ("SD_SCAN", "仅 IDLE", "忙时返回 SD_SCAN_BUSY，不会切换活动卡槽供电。"),
            ],
            [2520, 2700, 4140],
            code_columns=(0,),
            small=True,
        )

        self.heading("12.1 上位机推荐状态机", 2)
        self.numbered([
            "连接后以 STATUS 为设备状态真值，不只依赖按钮本地状态。",
            "启动任何模式前检查前置连接、SD 状态和授时状态，并锁定会冲突的控件。",
            "收到 STATUS=0 的启动响应后，再把界面切换为运行中；失败时展示完整 MSG。",
            "运行期间定时 STATUS，监控模式、DATA、REC、SLOT 和 SAMPLES。",
            "停止后再次 STATUS，确认回到 IDLE，再允许格式化、文件下载或换卡。",
        ])
        self.note(
            "不要并行发控制命令",
            "SEQ 由设备生成，当前协议没有客户端事务 ID。建议控制层串行化所有命令，每次只保留一个等待响应的请求，并设置 1～3 秒超时。",
        )

    def section_troubleshooting(self):
        self.heading("13 常见问题与定位", 1)
        self.table(
            ["现象", "优先检查", "处理建议"],
            [
                ("WinError 10054 / 连接后断开", "IP 冲突、设备复位、同端口已有客户端、W5500 socket 重建", "先 ping/检查网段，关闭旧客户端，分别重连 5001、5000，再发 STATUS。"),
                ("控制命令无响应", "无换行、大小写错误、参数名错误、命令超过 127 字节", "发送规范 ASCII + \\n；一发一收；打印发送字节和超时。"),
                ("START_REALTIME_FAIL", "STATUS 中 DATA 是否为 1", "先建立 5001 接收线程，再 START_RT。"),
                ("实时模式自动停止并报告 buffer overflow", "上位机读取速率长期不足", "修复接收线程，减少界面逐点处理和同步磁盘写入，重连后重新 START_RT。"),
                ("实际采样率与设置值不同", "包头 sample_rate 和设备 actual_clk", "这是定时器整数分频量化；分析时优先使用包头实际值。"),
                ("波形幅值不对", "字节序、24 位符号扩展、VREF、GAIN", "先打印前几个原始码值，再按本手册公式转换并做硬件标定。"),
                ("SD 显示 U", "卡存在但挂载/文件系统失败", "空闲时按需备份后格式化为 EXFAT，再 SD_SCAN。"),
                ("SD 显示 -", "未插卡或 HAL 初始化失败", "断电检查插槽和卡，再 SD_SCAN；当前网络响应不能细分原因。"),
                ("START_STORE 返回 SD 摘要失败", "writable=0", "先 SD_SCAN；至少一槽必须为 R。"),
                ("TASK_ADD_TIME_INVALID", "设备尚未授时或重启后授时失效", "重新 TIME，再添加任务。"),
                ("任务提前/延后执行", "电脑时区、Unix 毫秒、任务重叠", "核对本地时区和 START；读取 TASK_GET；避免重叠。"),
                ("TRIG_ARM STATUS=-1", "非 IDLE、无可写卡或数据端口未连接", "根据 OUTPUT 满足前置条件后重新布防。"),
                ("LIST/GET 返回 ERR BUSY", "设备不在 IDLE", "停止采集或取消触发布防，确认 IDLE 后重试。"),
                ("GET 下载卡住或文件偏短", "SD 读错或网络中断", "设置接收超时，以 SIZE 校验；不完整文件不得作为有效记录。"),
                ("实时 sequence 跳变", "上位机来不及收取、解析错位或链路丢数据", "记录序号、payload_len 和重同步次数；序号不会因正常重启流或连续触发分段而清零。"),
            ],
            [2520, 3060, 3780],
            small=True,
        )

    def section_acceptance(self):
        self.heading("14 上位机联调验收清单", 1)
        self.heading("14.1 基础通信", 2)
        self.bullets([
            "电脑和设备同网段，5000/5001/5002 可分别建立连接。",
            "STATUS 能正确解析全部字段；设备复位后上位机能发现断线并重连。",
            "未知或超时命令不会阻塞控制发送队列。",
        ])
        self.heading("14.2 实时链路", 2)
        self.bullets([
            "16/24 两种 PREC 均能按小端正确解码；payload_len 校验通过；界面不存在 32 bit 选项。",
            "1 kS/s、96 kS/s 实时传输包序号连续，128/200 kS/s 存储时实际样本率和写入统计合理。",
            "断开 5001 后上位机报警，设备自动停止采集并回到 IDLE。",
            "波形、原始码值、电压值和 CSV 保存互不阻塞网络接收线程。",
        ])
        self.heading("14.3 存储与文件", 2)
        self.bullets([
            "分别插入 SD1～SD4，SD_SCAN 的 R/U/- 与实物一致。",
            "SD_DEMO 在目标卡产生正确长度和模式的测试文件。",
            "单卡格式化和一键格式化全部卡均有危险确认、进行中状态和完成弹窗；ALL 能正确报告 formatted/skipped/failed。",
            "立即、任务、触发三种写卡目录前缀和 meta source 可区分。",
            "跨卡记录能按 segment 合并，总样本数和总字节数一致。",
            "LIST 能浏览目录；GET 的本地文件大小与 SIZE 一致。",
        ])
        self.heading("14.4 时间、任务和触发", 2)
        self.bullets([
            "年月日时分秒转换后的 Unix 毫秒与电脑时区一致。",
            "任务添加、查询、单删、全清、重复执行均通过，且不存在重叠时间窗。",
            "PC13 高电平时不触发；低电平稳定 10 ms 后触发；到时自动停止。",
            "触发 STORAGE 和 REALTIME 两种输出都能完成，结束后状态回到 IDLE。",
        ])

    def section_delivery_requirements(self):
        self.heading("15 委托开发交付要求", 1)
        self.paragraph(
            "本章是第三方上位机的最低交付边界。实现语言和 GUI 框架不限，但协议、线程模型、错误提示和验收行为必须满足下列要求。"
        )

        self.heading("15.1 必备界面", 2)
        self.table(
            ["界面/区域", "必须提供的能力"],
            [
                ("连接与设备概览", "IP 与三个端口、连接/断开、STATUS、设备模式、控制/数据连接、缓冲和丢样统计。"),
                ("实时采集", "独立 RATE/PREC/GAIN 配置、开始/停止、波形、样本率、包序号、丢包、原始码值、电压、Vpp、CSV 保存。"),
                ("立即自存储", "独立 RATE/PREC/GAIN 配置、开始/停止、当前卡槽、运行号、累计样本、切卡次数。"),
                ("任务表", "独立采集配置、年月日时分秒、持续时间、间隔、次数、添加、读取 16 项、单删、全清和表格展示。"),
                ("触发模式", "独立采集配置、STORAGE/REALTIME、ONCE/CONTINUOUS、持续时间、配置、布防、停止和状态查询。"),
                ("SD 管理", "SD_SCAN/SD_STATUS、四卡状态与容量、单卡格式化、全部卡格式化、SD_DEMO。"),
                ("SD 文件", "卡槽和目录浏览、逐级 LIST、文件 GET、进度/超时、SIZE 完整性校验。"),
                ("时间", "默认自动跟随电脑时间；用户手动编辑后停止自动刷新；发送前转换为 Unix 毫秒。"),
            ],
            [2340, 7020],
            small=True,
        )
        self.note(
            "交互要求",
            "界面全中文，四种采集模式分区明确。产品界面不得保留“手动命令”输入框。每个按钮都要有发送中状态；"
            "收到控制响应后根据 STATUS 弹出成功或失败提示，并保留完整 MSG 到日志。危险操作必须二次确认。",
            warning=True,
        )

        self.heading("15.2 推荐软件结构", 2)
        self.numbered([
            "控制线程维护 5000 的行缓冲和串行命令队列；同一时刻只允许一个待响应请求。",
            "数据线程专门阻塞读取 5001，使用 bytearray 累计并循环拆 HY 包，绝不在接收线程绘图或逐点写界面。",
            "文件线程按一次 LIST/GET 一条短连接处理 5002；GET 先读首行，再严格读取 SIZE 字节。",
            "GUI 线程通过有界队列批量取数，波形按固定刷新周期降采样绘制；CSV/文件落盘使用独立工作线程或批量缓冲。",
            "设备状态模型以 STATUS、SD_STATUS、TASK_GET、TRIG_STATUS 的设备返回为真值，不以本地按钮状态猜测。",
            "连接中断后清空待响应事务并提示用户；重连后依次 STATUS、SD_STATUS、TASK_LIST/TASK_GET、TRIG_STATUS，必要时重新 TIME。",
        ])

        self.heading("15.3 响应与超时策略", 2)
        self.table(
            ["场景", "建议策略"],
            [
                ("普通控制命令", "默认 3 s；解析 SEQ、STATUS 和完整 MSG。STATUS=0 才算成功。"),
                ("SD_SCAN / START_STORE", "可能含逐卡初始化或目录准备，建议 10～30 s，期间禁用冲突按钮。"),
                ("SD_FORMAT", "大容量卡可能较慢，建议至少 120 s，不得因界面超时自动重发。"),
                ("文件 LIST", "建议 5 s 空闲超时，收到 END LIST 才完成。"),
                ("文件 GET", "建议 8 s 空闲超时；总进度按 SIZE，短文件视为失败。"),
                ("无控制响应", "可能是命令语法被丢弃；记录原始发送字节，先 STATUS 核实，不自动重复破坏性命令。"),
            ],
            [2700, 6660],
            small=True,
        )

        self.heading("15.4 交付物", 2)
        self.bullets([
            "可安装运行程序、完整源代码、依赖锁定文件和构建说明。",
            "协议层、HY 拆包、24 位符号扩展、任务/SD 响应解析的自动化测试。",
            "异常断线、控制超时、数据端慢读、文件下载中断和格式化失败的测试记录。",
            "本手册第 14 章逐项签字的联调报告，包含设备固件版本、上位机版本、测试卡型号和网络环境。",
        ])

    def appendix_commands(self):
        heading = self.heading("附录 A 控制命令速查", 1)
        heading.paragraph_format.page_break_before = True
        rows = [
            ("STATUS", "任意", "设备模式和关键状态"),
            ("TIME <unix_ms>", "任意", "TIME_OK"),
            ("CONFIG RATE=<Hz> PREC=<16|24> GAIN=<dB>", "任意；下次启动生效", "CONFIG_OK / CONFIG_INVALID"),
            ("START_RT", "5001 已连接", "START_REALTIME / START_REALTIME_FAIL"),
            ("STOP_RT", "任意", "STOP_REALTIME"),
            ("START_STORE", "至少一张 R 卡", "START_STORAGE / SD 摘要"),
            ("STOP_STORE", "任意", "STOP_STORAGE"),
            ("SD_STATUS", "任意", "缓存的 SD 摘要"),
            ("SD_SCAN", "仅 IDLE", "刷新后的 SD 摘要 / SD_SCAN_BUSY"),
            ("SD_FORMAT SLOT=n FS=EXFAT CONFIRM=YES", "仅 IDLE", "SD_FORMAT_OK / *_FAIL / *_BUSY"),
            ("SD_FORMAT SLOT=ALL FS=EXFAT CONFIRM=ERASE_ALL", "仅 IDLE；危险操作", "SD_FORMAT_ALL_OK / PARTIAL / NO_CARD"),
            ("SD_DEMO SLOT=ALL BYTES=4096", "仅 IDLE", "SD_DEMO tried=..."),
            ("TASK_ADD ...", "已授时", "TASK_ADD / TASK_ADD_TIME_INVALID"),
            ("TASK_LIST", "任意", "TASK_LIST COUNT=n MAX=16 TIME=0|1"),
            ("TASK_GET INDEX=n", "n=0～15", "TASK ... / TASK_EMPTY / TASK_GET_INVALID"),
            ("TASK_DELETE INDEX=n", "n=0～15 且非空", "TASK_DELETE INDEX=n，状态码判成功"),
            ("TASK_CLEAR", "任意", "TASK_CLEAR"),
            ("TRIG_CONFIG OUTPUT=<STORAGE|REALTIME> REPEAT=<ONCE|CONTINUOUS> ...", "仅 IDLE", "TRIG_CONFIG_OK / TRIG_CONFIG_BUSY"),
            ("TRIG_ARM", "仅 IDLE 且满足输出前置条件", "TRIG_ARM level=LOW，状态码判成功"),
            ("TRIG_STATUS", "任意", "TRIG_STATUS ..."),
            ("TRIG_STOP", "任意", "TRIG_STOP"),
        ]
        self.table(["命令", "前置条件", "响应 MSG"], rows, [4500, 2160, 2700], code_columns=(0,), small=True)

        self.heading("附录 B 文件端口速查", 1)
        self.table(
            ["命令", "首行响应", "后续内容"],
            [
                ("LIST SLOT=n PATH=/...", "OK LIST ...", "0～64 条 ENTRY，最后 END LIST COUNT=n"),
                ("GET SLOT=n PATH=/...", "OK FILE ... SIZE=n", "紧跟恰好 n 字节二进制文件，无尾标记"),
            ],
            [3600, 3060, 2700],
            code_columns=(0, 1),
        )

    def appendix_limits(self):
        self.heading("附录 C 当前版本约束", 1)
        self.table(
            ["项目", "当前实现", "上位机策略"],
            [
                ("配置、授时、任务持久化", "全部为 RAM，复位/掉电丢失", "每次上线重新授时、配置和下发任务"),
                ("控制并发", "无客户端事务 ID，设备生成 SEQ", "命令串行，一发一收"),
                ("任务重叠", "无排队/优先级仲裁", "下发前校验时间窗"),
                ("实时断线/溢出", "设备自动停止实时采集", "连接监控，恢复后重新 START_RT"),
                ("实时时间戳", "包内无采样时间", "按实际 rate 推算；电脑时间仅作收包时间"),
                ("应用层完整性", "无 CRC/哈希", "校验长度和 sequence；重要文件可在上位机生成哈希"),
                ("文件列表", "最多 64 项，无分页", "目录分层，保留已知路径索引"),
                ("文件操作", "只读 LIST/GET", "删除和整理通过拔卡读卡器完成"),
                ("记录实际采样率", "包头与 meta 均保存 actual rate", "离线时间轴使用 actual_sample_rate_hz"),
                ("安全", "明文 TCP、无认证", "仅在受控网络使用"),
                ("最大容量", "代码模型 4×2 TiB", "用目标品牌/容量实卡做格式化、满盘、跨卡压力测试"),
            ],
            [2340, 3060, 3960],
            small=True,
        )

        self.heading("附录 D 源码基线", 1)
        self.paragraph("本手册主要按以下当前源码核对生成：")
        self.bullets([
            "device/command_protocol.c：控制命令解析和响应格式。",
            "device/app_controller.c：状态机、命令行为、触发和模式切换。",
            "device/tcp_service.c 与 stream_service.c：端口、实时包头和发送逻辑。",
            "device/storage_manager.c 与 storage_recorder.c：卡槽、格式化、目录、meta.txt 和 data.bin。",
            "device/task_scheduler.c 与 time_service.c：任务重复规则和软件时钟。",
            "device/file_service.c：LIST/GET 文件协议。",
            "tools/host_gui/hydrophone_gui.py：当前 GUI 的时间转换、数据解码和电压换算。",
        ])

    def save(self, output):
        output.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("output") / "单通道自容式水听器上位机对接手册.docx",
    )
    args = parser.parse_args()

    builder = ManualBuilder()
    builder.cover()
    builder.front_matter()
    builder.build_content()
    builder.save(args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
